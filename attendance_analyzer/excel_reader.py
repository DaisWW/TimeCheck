from __future__ import annotations

from collections import defaultdict
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from .approval import LEAVE_KEYWORDS, approval_parse_issues
from .config import AppConfig
from .models import AttendanceData, DailyAttendance, InputIssue, PunchRecord
from .time_utils import combine_date_time, parse_date_text, parse_datetime, parse_shift_times


DAILY_FIELDS = (
    "姓名", "UserId", "日期", "班次", "上班1打卡时间", "上班1打卡结果",
    "下班1打卡时间", "下班1打卡结果", "关联的审批单",
)
RAW_FIELDS = ("姓名", "UserId", "考勤日期", "打卡时间", "打卡结果")
PUNCH_RESULTS = frozenset((
    "正常", "迟到", "早退", "严重迟到", "旷工迟到", "缺卡", "未打卡", "补卡审批通过", *LEAVE_KEYWORDS,
))


class WorkbookInputError(ValueError):
    def __init__(self, path: Path, issues: list[InputIssue]):
        self.issues = issues
        lines = [f"输入校验失败：{path}（{len(issues)} 项），未生成统计结果。"]
        for issue in issues:
            lines.append(
                f"{issue.source_sheet} 第 {issue.source_row} 行 [{issue.field_name}] "
                f"原始值={issue.raw_value!r}：{issue.message}"
            )
        super().__init__("\n".join(lines))


class DingTalkWorkbookReader:
    def __init__(self, config: AppConfig):
        self._config = config

    def read(self, path: Path) -> AttendanceData:
        # 不使用公式缓存，避免公式或未缓存的结果被当作空白输入。
        try:
            workbook = load_workbook(path, data_only=False, read_only=False)
        except KeyError as error:
            raise ValueError(f"{path}：不是完整的 Excel 工作簿，请从钉钉重新导出原始考勤后重试。") from error
        try:
            for name in (self._config.sheets.daily_sheet, self._config.sheets.raw_sheet):
                if name not in workbook.sheetnames:
                    raise ValueError(f"{path}：缺少工作表 {name!r}；实际工作表：{workbook.sheetnames}")
            daily_sheet = workbook[self._config.sheets.daily_sheet]
            raw_sheet = workbook[self._config.sheets.raw_sheet]
            input_issues: list[InputIssue] = []
            _validate_headers(daily_sheet, (3, 4), DAILY_FIELDS, input_issues)
            _validate_headers(raw_sheet, (3,), RAW_FIELDS, input_issues)
            if input_issues:
                raise WorkbookInputError(path, input_issues)
            daily_records = self._read_daily_records(daily_sheet, input_issues)
            raw_punches = self._read_raw_punches(raw_sheet, input_issues)
            self._validate_relationships(daily_sheet, raw_sheet, daily_records, raw_punches, input_issues)
            blocking = [issue for issue in input_issues if issue.blocks_summary]
            if blocking:
                raise WorkbookInputError(path, blocking)
            if not daily_records:
                raise ValueError(f"{path}：每日统计没有考勤数据，不能生成空统计结果。")
            return AttendanceData(daily_records, raw_punches, input_issues)
        finally:
            workbook.close()

    def _read_daily_records(
        self, sheet: Worksheet, input_issues: list[InputIssue] | None = None
    ) -> list[DailyAttendance]:
        # 钉钉“每日统计”是双层表头：第 3 行主标题，第 4 行子标题。
        headers = _headers(sheet, header_rows=(3, 4))
        rows: list[DailyAttendance] = []
        input_issues = input_issues if input_issues is not None else []
        max_column = sheet.max_column
        for row in range(5, sheet.max_row + 1):
            if not _row_has_data(sheet, row, max_column):
                continue
            name = _cell(sheet, row, headers, "姓名")
            date_text = _cell(sheet, row, headers, "日期")
            work_date = parse_date_text(_value(sheet, row, headers, "日期"))
            if not name:
                if _row_has_data(sheet, row, max_column):
                    input_issues.append(
                        _input_issue(
                            sheet,
                            row,
                            "",
                            "",
                            None,
                            "姓名",
                            name,
                            "每日统计缺少姓名，无法归属该考勤记录。",
                            "该行不参与统计。",
                            "补全姓名后重新导出并运行。",
                        )
                    )
                continue
            if work_date is None:
                input_issues.append(
                    _input_issue(
                        sheet,
                        row,
                        name,
                        _cell(sheet, row, headers, "UserId"),
                        None,
                        "日期",
                        date_text,
                        "每日统计日期为空或无法解析。",
                        "该行不参与统计。",
                        "修正日期后重新导出并运行。",
                    )
                )
                continue

            user_id = _cell(sheet, row, headers, "UserId")
            start_time_text = _cell(sheet, row, headers, "上班1打卡时间")
            end_time_text = _cell(sheet, row, headers, "下班1打卡时间")
            approval_text = _cell(sheet, row, headers, "关联的审批单")
            self._validate_row_fields(sheet, row, headers, DAILY_FIELDS, name, user_id, work_date, input_issues)
            for field_name, time_text in (
                ("上班1打卡时间", start_time_text),
                ("下班1打卡时间", end_time_text),
            ):
                if time_text and combine_date_time(work_date, time_text) is None:
                    input_issues.append(
                        _input_issue(
                            sheet,
                            row,
                            name,
                            user_id,
                            work_date,
                            field_name,
                            time_text,
                            "每日统计打卡时间无法解析。",
                            "每日统计打卡兜底不可用；可能导致缺卡或加班不计入。",
                            "修正打卡时间后重新导出并运行。",
                            issue_type="解析问题",
                        )
                    )
            approval_issues = approval_parse_issues(approval_text, work_date.year)
            if approval_issues:
                messages = list(dict.fromkeys(issue.message for issue in approval_issues))
                blocks_summary = any(issue.blocks_summary for issue in approval_issues)
                input_issues.append(
                    _input_issue(
                        sheet,
                        row,
                        name,
                        user_id,
                        work_date,
                        "关联的审批单",
                        approval_text,
                        "；".join(messages),
                        "加班审批核对或请假抵扣无法自动执行。",
                        "在钉钉审批中确认起止时间后重新导出并运行。",
                        issue_type="解析问题" if blocks_summary else "业务待确认",
                        blocks_summary=blocks_summary,
                    )
                )

            additional_punches = []
            for prefix in ("上班2", "下班2", "上班3", "下班3"):
                time_field, result_field = prefix + "打卡时间", prefix + "打卡结果"
                if time_field not in headers:
                    continue
                value = _cell(sheet, row, headers, time_field)
                result_value = _cell(sheet, row, headers, result_field)
                self._validate_row_fields(sheet, row, headers, (time_field,), name, user_id, work_date, input_issues)
                if result_field in headers:
                    self._validate_row_fields(sheet, row, headers, (result_field,), name, user_id, work_date, input_issues)
                timestamp = combine_date_time(work_date, value)
                if value and (timestamp is None or (timestamp.date() - work_date).days not in (0, 1)):
                    input_issues.append(_input_issue(
                        sheet, row, name, user_id, work_date, time_field, value,
                        "额外打卡时间无法解析或不属于考勤当天/次日。", "无法确认首末有效打卡。", "核对该打卡时间。",
                    ))
                additional_punches.append((value, result_value))
            first = combine_date_time(work_date, start_time_text)
            last = combine_date_time(work_date, end_time_text)
            if ((first and first.date() != work_date)
                or (last and (last.date() - work_date).days not in (0, 1))
                or (first and last and last <= first)):
                input_issues.append(_input_issue(
                    sheet, row, name, user_id, work_date, "上下班打卡时间",
                    f"{start_time_text} → {end_time_text}",
                    "上下班日期或先后顺序异常；跨午夜下班请明确写次日，不自动交换上下班。",
                    "无法自动计算。", "核对日期及上下班方向。",
                ))
            shift = _cell(sheet, row, headers, "班次")
            if not shift and (start_time_text or end_time_text or approval_text or any(value for value, _ in additional_punches)):
                input_issues.append(_input_issue(
                    sheet, row, name, user_id, work_date, "班次", shift,
                    "有打卡或审批但缺少班次，不能默认按工作日计算。",
                    "无法确认日期类型或工作时段。", "补齐当天排班。",
                ))

            rows.append(
                DailyAttendance(
                    employee_name=name,
                    user_id=user_id,
                    work_date=work_date,
                    shift=_cell(sheet, row, headers, "班次"),
                    start_time_text=start_time_text,
                    start_result=_cell(sheet, row, headers, "上班1打卡结果"),
                    end_time_text=end_time_text,
                    end_result=_cell(sheet, row, headers, "下班1打卡结果"),
                    approval_text=approval_text,
                    source_row=row,
                    additional_punches=tuple(additional_punches),
                )
            )
        return rows

    def _read_raw_punches(
        self, sheet: Worksheet, input_issues: list[InputIssue] | None = None
    ) -> dict[tuple[str, date], list[PunchRecord]]:
        # 钉钉“原始记录”只有第 3 行是表头，第 4 行开始就是数据。
        headers = _headers(sheet, header_rows=(3,))
        grouped: dict[tuple[str, date], list[PunchRecord]] = defaultdict(list)
        input_issues = input_issues if input_issues is not None else []
        max_column = sheet.max_column

        for row in range(4, sheet.max_row + 1):
            if not _row_has_data(sheet, row, max_column):
                continue
            name = _cell(sheet, row, headers, "姓名")
            user_id = _cell(sheet, row, headers, "UserId")
            occurred_at_text = _cell(sheet, row, headers, "打卡时间")
            attendance_date_text = _cell(sheet, row, headers, "考勤日期")
            occurred_at = parse_datetime(_value(sheet, row, headers, "打卡时间"))
            attendance_date = parse_date_text(_value(sheet, row, headers, "考勤日期"))
            result = _cell(sheet, row, headers, "打卡结果")
            if not name:
                if _row_has_data(sheet, row, max_column):
                    input_issues.append(
                        _input_issue(
                            sheet,
                            row,
                            "",
                            user_id,
                            attendance_date,
                            "姓名",
                            name,
                            "原始记录缺少姓名，无法归属该打卡。",
                            "该打卡不参与统计。",
                            "补全姓名后重新导出并运行。",
                        )
                    )
                continue
            self._validate_row_fields(sheet, row, headers, RAW_FIELDS, name, user_id, attendance_date, input_issues)
            if occurred_at is None:
                input_issues.append(
                    _input_issue(
                        sheet,
                        row,
                        name,
                        user_id,
                        attendance_date,
                        "打卡时间",
                        occurred_at_text,
                        "原始打卡时间为空或无法解析。",
                        "该打卡不参与统计；可能导致缺卡或加班不计入。",
                        "修正打卡时间后重新导出并运行。",
                    )
                )
            if attendance_date is None:
                input_issues.append(
                    _input_issue(
                        sheet,
                        row,
                        name,
                        user_id,
                        None,
                        "考勤日期",
                        attendance_date_text,
                        "原始考勤日期为空或无法解析，无法确认跨午夜归属。",
                        "该打卡不参与统计。",
                        "修正考勤日期后重新导出并运行。",
                        issue_type="解析问题",
                    )
                )
            if occurred_at is None or attendance_date is None:
                continue

            if (occurred_at.date() - attendance_date).days not in (0, 1):
                input_issues.append(_input_issue(
                    sheet, row, name, user_id, attendance_date, "打卡时间", occurred_at_text,
                    "打卡时间必须属于考勤当天或次日；不接受提前或跨多日的打卡归属。",
                    "无法自动计算。", "核对打卡时间和考勤日期。",
                ))
            if any(result.startswith(prefix) for prefix in self._config.overtime.invalid_punch_prefixes):
                input_issues.append(
                    _input_issue(
                        sheet,
                        row,
                        name,
                        user_id,
                        attendance_date,
                        "打卡结果",
                        result,
                        "原始记录明确标记为打卡无效，当前不作为有效打卡。",
                        "该行不参与首末有效打卡；可能导致成对打卡不足。",
                        "确认该打卡是否应恢复有效；修正源数据后重新运行。",
                        issue_type="原始数据被规则排除",
                        blocks_summary=False,
                    )
                )

            # 跨午夜的下班卡在“打卡时间”属于次日，但钉钉的“考勤日期”明确归属前一考勤日。
            key = (_person_key(user_id, name), attendance_date)
            grouped[key].append(
                PunchRecord(
                    employee_name=name,
                    user_id=user_id,
                    occurred_at=occurred_at,
                    result=result,
                    source_row=row,
                )
            )

        for punches in grouped.values():
            punches.sort(key=lambda punch: punch.occurred_at)
        return dict(grouped)

    def _validate_row_fields(self, sheet, row, headers, fields, name, user_id, work_date, issues):
        for field in fields:
            cell = sheet.cell(row, headers[field])
            value = _cell(sheet, row, headers, field)
            message = ""
            if cell.data_type in {"f", "e"}:
                message = "输入字段含公式或 Excel 错误值，必须提供明确的原始值。"
            elif field == "UserId" and cell.value is not None and not isinstance(cell.value, str):
                message = "UserId 必须以文本保存，避免长数字精度丢失。"
            elif field == "班次" and value and not any(keyword in value for keyword in self._config.workday.rest_shift_keywords):
                try:
                    parse_shift_times(value)
                except ValueError as error:
                    message = str(error)
            elif "打卡结果" in field and value not in PUNCH_RESULTS:
                invalid = any(value.startswith(prefix) for prefix in self._config.overtime.invalid_punch_prefixes)
                if not invalid and (value or fields == RAW_FIELDS):
                    message = "未知或空的打卡结果，不能默认作为有效打卡。"
            if message:
                issues.append(_input_issue(sheet, row, name, user_id, work_date, field, value,
                                           message, "无法自动计算。", "核对源单元格或确认新增格式。"))

    def _validate_relationships(self, daily_sheet, raw_sheet, days, grouped, issues):
        keys = {}
        employee_names = {}
        for day in days:
            employee_key = person_key(day.user_id, day.employee_name)
            key = (employee_key, day.work_date)
            if employee_key in employee_names and employee_names[employee_key] != day.employee_name:
                issues.append(_input_issue(
                    daily_sheet, day.source_row, day.employee_name, day.user_id, day.work_date,
                    "姓名/UserId", f"{day.employee_name}/{day.user_id}", "同一身份键对应不同姓名，不能合并考勤。",
                    "可能混合员工记录。", "核对姓名及 UserId。",
                ))
            employee_names[employee_key] = day.employee_name
            if key in keys:
                issues.append(_input_issue(
                    daily_sheet, day.source_row, day.employee_name, day.user_id, day.work_date,
                    "姓名/UserId/日期", str(key), f"同一员工同一天重复，首条在第 {keys[key]} 行。",
                    "会重复统计并消耗迟到宽限。", "确认并移除重复记录。",
                ))
            keys[key] = day.source_row
        for key, punches in grouped.items():
            for punch in punches:
                if key[0] in employee_names and punch.employee_name != employee_names[key[0]]:
                    issues.append(_input_issue(
                        raw_sheet, punch.source_row, punch.employee_name, punch.user_id, key[1],
                        "姓名/UserId", f"{punch.employee_name}/{punch.user_id}", "两表相同身份键的姓名不一致。",
                        "无法确认员工身份。", "核对两表姓名及 UserId。",
                    ))
                if key not in keys:
                    issues.append(_input_issue(
                        raw_sheet, punch.source_row, punch.employee_name, punch.user_id, key[1],
                        "姓名/UserId/考勤日期", str(key), "原始打卡找不到对应的每日统计，不能静默丢弃。",
                        "无法确认打卡归属。", "核对两表身份和日期，补齐每日统计。",
                    ))


def _validate_headers(sheet, header_rows, required, issues):
    for field in required:
        columns = [column for column in range(1, sheet.max_column + 1)
                   if any(str(_merged_value(sheet, row, column) or "").strip() == field for row in header_rows)]
        if len(columns) != 1:
            issues.append(_input_issue(
                sheet, header_rows[0], "", "", None, field, str(columns),
                "缺少必需表头或存在重复表头；不允许当作空列继续计算。",
                "无法解析工作表结构。", "使用钉钉每日统计/原始记录的完整导出格式。",
            ))


def _headers(sheet: Worksheet, header_rows: tuple[int, ...]) -> dict[str, int]:
    result: dict[str, int] = {}
    for column in range(1, sheet.max_column + 1):
        parts: list[str] = []
        for row in header_rows:
            value = _merged_value(sheet, row, column)
            if value is None:
                continue
            text = str(value).strip()
            if text and text not in parts:
                parts.append(text)
        if parts:
            result[" - ".join(parts)] = column
            result.setdefault(parts[-1], column)
    return result


def _merged_value(sheet: Worksheet, row: int, column: int) -> Any:
    value = sheet.cell(row, column).value
    if value is not None:
        return value
    for cell_range in sheet.merged_cells.ranges:
        if cell_range.min_row <= row <= cell_range.max_row and cell_range.min_col <= column <= cell_range.max_col:
            return sheet.cell(cell_range.min_row, cell_range.min_col).value
    return None


def _cell(sheet: Worksheet, row: int, headers: dict[str, int], name: str) -> str:
    value = _value(sheet, row, headers, name)
    return "" if value is None else str(value).strip()


def _value(sheet: Worksheet, row: int, headers: dict[str, int], name: str) -> Any:
    column = headers.get(name)
    if column is None:
        return None
    return sheet.cell(row, column).value


def _row_has_data(sheet: Worksheet, row: int, max_column: int) -> bool:
    return any(
        value is not None and str(value).strip()
        for value in (sheet.cell(row, column).value for column in range(1, max_column + 1))
    )


def _input_issue(
    sheet: Worksheet,
    source_row: int,
    employee_name: str,
    user_id: str,
    work_date: date | None,
    field_name: str,
    raw_value: str,
    message: str,
    impact: str,
    suggested_action: str,
    issue_type: str = "解析问题",
    blocks_summary: bool = True,
) -> InputIssue:
    return InputIssue(
        source_sheet=sheet.title,
        source_row=source_row,
        employee_name=employee_name,
        user_id=user_id,
        work_date=work_date,
        field_name=field_name,
        raw_value=raw_value,
        message=message,
        impact=impact,
        suggested_action=suggested_action,
        issue_type=issue_type,
        blocks_summary=blocks_summary,
    )


def _person_key(user_id: str, name: str) -> str:
    return user_id.strip() or name.strip()


def fallback_punches(day: DailyAttendance) -> list[PunchRecord]:
    punches: list[PunchRecord] = []
    for value, result in (
        (day.start_time_text, day.start_result), (day.end_time_text, day.end_result), *day.additional_punches
    ):
        timestamp = combine_date_time(day.work_date, value)
        if timestamp is None:
            continue
        punches.append(
            PunchRecord(
                employee_name=day.employee_name,
                user_id=day.user_id,
                occurred_at=timestamp,
                result=result,
                source_row=day.source_row,
            )
        )
    punches.sort(key=lambda punch: punch.occurred_at)
    return punches


def person_key(user_id: str, name: str) -> str:
    return _person_key(user_id, name)
