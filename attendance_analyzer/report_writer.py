from __future__ import annotations

from collections import Counter, defaultdict
from collections.abc import Callable
from copy import copy
from datetime import date, datetime, timedelta
from pathlib import Path
import re
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook
from openpyxl.cell import Cell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.writer.excel import ExcelWriter

from .config import AppConfig
from .models import InputIssue, OvertimeResult
from .rules import NON_DEDUCTIBLE_LEAVE_TYPES
from .time_utils import fmt_dt, fmt_time


KNOWN_LEAVE_TYPES = ("年假", "调休", "事假", "病假", "婚假", "外出", "出差", "请假")
INVALID_SHEET_CHARS = re.compile(r"[\\/*?:\[\]]")

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
RED_FILL = PatternFill("solid", fgColor="FFC7CE")
RED_FONT = Font(color="9C0006", bold=True)
AMBER_FILL = PatternFill("solid", fgColor="FFEB9C")
AMBER_FONT = Font(color="9C6500")
YELLOW_FILL = PatternFill("solid", fgColor="FFF2CC")
GRID_BORDER = Border(bottom=Side(style="hair", color="D9E2F3"))


class OvertimeReportWriter:
    """Write one auditable workbook for one input month."""

    def __init__(self, config: AppConfig, *, stage: Callable[[str, str], None] | None = None):
        self._config = config
        self._stage = stage

    def write(
        self,
        output_path: Path,
        results: list[OvertimeResult],
        input_issues: list[InputIssue] | None = None,
    ) -> None:
        input_issues = input_issues or []
        results = sorted(results, key=lambda item: (item.employee_name, item.work_date, item.user_id))
        grouped = _group_results(results)
        period_start, period_end = _period_bounds(results)

        workbook = Workbook()
        if self._stage:
            self._stage("report", "正在生成总览和员工明细")
        overview = workbook.active
        overview.title = "总览"
        self._write_overview(overview, grouped, input_issues, period_start, period_end)

        # Employee tabs are the primary detailed view. They show every
        # calendar date in the input period, including days without punches.
        used_names = {"总览"}
        for employee_index, ((name, user_id), items) in enumerate(sorted(grouped.items()), 1):
            sheet_name = _unique_sheet_name(name, used_names, employee_index)
            employee_sheet = workbook.create_sheet(sheet_name)
            self._write_employee_detail(
                employee_sheet,
                name,
                user_id,
                items,
                input_issues,
                period_start,
                period_end,
            )
            if self._stage:
                self._stage("report", f"已生成 {employee_index} / {len(grouped)} 份员工明细")

        self._write_supplement_punch_notifications(workbook.create_sheet("待补卡通知"), results)
        self._write_approval_reviews(workbook.create_sheet("待审核清单"), results)
        self._write_input_issues(workbook.create_sheet("待人工确认"), input_issues)
        self._write_anomalies(workbook.create_sheet("异常数据"), results, input_issues)
        self._write_daily_summary(workbook.create_sheet("每日汇总"), results, input_issues)

        # Retain the former all-person tabs for existing HR scripts.
        self._write_detail(workbook.create_sheet("加班明细"), results)
        self._write_employee_summary(workbook.create_sheet("员工汇总"), results, input_issues)
        self._write_config(workbook.create_sheet("配置"), output_path, input_issues)

        for sheet in workbook.worksheets:
            _finish_sheet(sheet)

        output_path.parent.mkdir(parents=True, exist_ok=True)
        if self._stage:
            self._stage("save", "正在保存 Excel 报表，请稍候")
        with ZipFile(output_path, "w", ZIP_DEFLATED, allowZip64=True, compresslevel=1) as archive:
            ExcelWriter(workbook, archive).save()

    def _write_overview(
        self,
        sheet,
        grouped: dict[tuple[str, str], list[OvertimeResult]],
        input_issues: list[InputIssue],
        period_start: date | None,
        period_end: date | None,
    ) -> None:
        leave_types = _leave_types(grouped.values())
        legacy_headers = [
            "姓名", "UserId", "统计期间", "统计状态", "记录天数", "有效统计天数",
            "加班天数", "加班总分钟", "加班总小时", "工作日加班分钟", "休息日加班分钟",
            "节假日加班分钟", "待审核加班分钟", "漏卡天数", "请假审批覆盖分钟(原始)",
            "请假有效分钟(扣午休)",
        ]
        legacy_headers.extend(
            f"{kind}分钟" + ("(不扣)" if kind in NON_DEDUCTIBLE_LEAVE_TYPES else "")
            for kind in leave_types
        )
        legacy_headers.extend([
            "其他请假分钟", "不扣工资请假分钟", "扣工资请假分钟", "请假重叠分钟",
            "原始迟到次数", "原始迟到分钟", "请假抵扣迟到分钟", "规则后迟到次数",
            "规则后迟到分钟", "宽限迟到分钟",
            "扣工资迟到次数", "扣工资迟到分钟", "待审核天数", "异常天数",
            "人工确认条数", "阻断人工确认条数", "人工确认状态", "请假与加班并存天数",
        ])
        friendly_headers = [
            "加班时间（分钟）", "加班时间（小时）",
            "工作日加班时间（分钟）", "休息日加班时间（分钟）", "节假日加班时间（分钟）",
            "待审核加班时间（分钟）",
            "请假时间（分钟）", "请假时间（调休、年假）（分钟）", "请假时间（其他）（分钟）",
            "请假重叠时间（分钟）",
            "迟到次数（原始）", "迟到时间（分钟）", "请假抵扣时间（分钟）",
            "迟到次数（规则后）", "迟到宽限时间（分钟）", "迟到计入次数", "迟到计入时间（分钟）",
            "时间差额（分钟）",
            "待审核天数", "漏卡天数", "异常天数", "人工确认条数", "阻断人工确认条数",
            "人工确认状态", "请假与加班并存天数",
            "外出/出差出勤时间（分钟）", "工资分类待确认请假（分钟）",
            "公司放假加班时间（分钟）", "公司放假抵扣迟到时间（分钟）",
        ]
        headers = legacy_headers[:7] + friendly_headers + legacy_headers[7:]
        sheet.append(headers)

        # Keep the first occurrence when a friendly field intentionally mirrors
        # an older audit field.  The legacy columns remain available (hidden)
        # for existing HR scripts and backwards-compatible checks.
        header_index: dict[str, int] = {}
        for index, value in enumerate(headers):
            header_index.setdefault(str(value), index)

        period_text = _period_text(period_start, period_end)
        for (name, user_id), items in sorted(grouped.items()):
            all_issue_count, blocking_issue_count = _person_issue_counts(
                input_issues, name, user_id
            )
            row = _overview_row(
                name, user_id, period_text, items, all_issue_count,
                blocking_issue_count, self._config, leave_types,
            )
            friendly_row = _overview_friendly_row(
                items,
                self._config,
                all_issue_count,
                blocking_issue_count,
            )
            # _overview_row retains the former, fully auditable column set.
            # Put the compact HR-facing block before those hidden audit fields.
            output_row = row[:7] + friendly_row + row[7:]
            cells = _append_row(sheet, output_row)
            if output_row[header_index["漏卡天数"]]:
                _mark_row(cells, RED_FILL, RED_FONT)
            elif output_row[header_index["待审核天数"]]:
                _mark_cells(
                    cells,
                    {header_index["统计状态"], header_index["待审核天数"]},
                    AMBER_FILL,
                    AMBER_FONT,
                )
            elif all_issue_count:
                _mark_cells(
                    cells,
                    {
                        header_index["统计状态"],
                        header_index["人工确认条数"],
                        header_index["阻断人工确认条数"],
                        header_index["人工确认状态"],
                    },
                    YELLOW_FILL,
                    None,
                )

        _style_header(sheet)
        sheet.auto_filter.ref = f"A1:{get_column_letter(sheet.max_column)}{sheet.max_row}"
        # The compact block above is the reader-facing view.  Keep the old
        # metric columns hidden but present for audit/backward compatibility.
        visible_columns = 7 + len(friendly_headers)
        for column in range(visible_columns + 1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(column)].hidden = True

    def _write_employee_detail(
        self,
        sheet,
        employee_name: str,
        user_id: str,
        items: list[OvertimeResult],
        input_issues: list[InputIssue],
        period_start: date | None,
        period_end: date | None,
    ) -> None:
        leave_types = _leave_types([items])
        headers = [
            "日期", "星期", "班次", "日类型", "每日统计上班打卡", "上班打卡结果",
            "每日统计下班打卡", "下班打卡结果", "首个有效打卡", "最后有效打卡",
            "成对有效打卡", "跨午夜按考勤日期归属", "原始记录数", "有效记录数",
            "使用每日统计兜底", "每日统计源行", "加班开始", "加班结束", "原始加班分钟",
            "午休扣除分钟", "晚餐扣除分钟", "计入加班分钟", "计入加班小时", "加班计算说明",
            "关联审批单", "请假类型", "请假审批覆盖分钟(原始)", "请假有效分钟(扣午休)",
        ]
        headers.extend(f"{kind}分钟" for kind in leave_types)
        headers.extend([
            "其他请假分钟", "不扣工资请假分钟", "扣工资请假分钟", "请假重叠分钟",
            "原始迟到分钟", "请假抵扣迟到分钟", "规则后迟到分钟", "迟到次数(月内)",
            "宽限迟到分钟", "扣工资迟到分钟", "迟到状态", "审批加班分钟", "审批差异分钟",
            "审批核对状态", "漏卡", "人工确认条数", "人工确认状态", "异常状态", "统计状态", "备注",
            "外出/出差出勤分钟", "工资分类待确认请假分钟",
            "公司放假时段", "公司放假说明", "公司放假加班分钟", "公司放假加班开始", "公司放假加班结束",
            "工作日加班分钟", "休息日加班分钟", "公司放假抵扣迟到分钟", "当天应出勤分钟",
        ])
        sheet.append(headers)

        header_index = {str(value): index for index, value in enumerate(headers)}

        by_date: dict[date, OvertimeResult] = {}
        for item in sorted(items, key=lambda value: value.work_date):
            by_date.setdefault(item.work_date, item)
        dates = sorted(by_date) if period_start is None or period_end is None else _date_range(period_start, period_end)

        for work_date in dates:
            result = by_date.get(work_date)
            day_issues = _issues_for_day(input_issues, employee_name, user_id, work_date)
            values = _employee_detail_row(result, work_date, day_issues, self._config, leave_types)
            holiday = self._config.company_holiday_on(work_date)
            cells = _append_row(sheet, values + [
                result.business_attendance_minutes if result else 0,
                result.leave_unclassified_minutes if result else 0,
                holiday.period if holiday else "", holiday.note if holiday else "",
                result.company_holiday_overtime_minutes if result else 0,
                fmt_dt(result.company_holiday_overtime_start) if result else "",
                fmt_dt(result.company_holiday_overtime_end) if result else "",
                result.workday_overtime_minutes if result else 0,
                result.rest_day_overtime_minutes if result else 0,
                result.late_company_holiday_exempt_minutes if result else 0,
                result.required_work_minutes if result else 0,
            ])
            if values[header_index["漏卡"]]:
                _mark_row(cells, RED_FILL, RED_FONT)
            elif values[header_index["人工确认状态"]] == "待处理":
                _mark_cells(
                    cells,
                    {header_index["人工确认状态"], header_index["异常状态"]},
                    YELLOW_FILL,
                    None,
                )
            elif (
                "待审核" in str(values[header_index["异常状态"]])
                or "待审核" in str(values[header_index["统计状态"]])
            ):
                _mark_cells(
                    cells,
                    {header_index["异常状态"], header_index["统计状态"]},
                    AMBER_FILL,
                    AMBER_FONT,
                )

        _style_header(sheet)
        sheet.auto_filter.ref = f"A1:{get_column_letter(sheet.max_column)}{sheet.max_row}"

    def _write_detail(self, sheet, results: list[OvertimeResult]) -> None:
        headers = [
            "姓名", "UserId", "日期", "班次", "加班类型", "上班打卡", "下班打卡", "弹性下班",
            "午休扣除(分钟)", "晚餐扣除(分钟)", "加班开始", "加班结束", "原始加班分钟",
            "计入加班分钟", "计入加班小时", "关联审批单", "请假审批覆盖当日分钟(非计薪)",
            "请假有效分钟(扣午休)", "请假类型", "不扣工资请假分钟", "扣工资请假分钟", "请假重叠分钟",
            "审批加班分钟", "审批差异(分钟)", "审批核对状态", "原始迟到分钟", "请假抵扣迟到分钟",
            "迟到分钟", "迟到次数(月内)", "宽限迟到分钟", "扣款迟到分钟", "迟到状态", "统计状态",
            "备注", "计算说明",
            "公司放假时段", "公司放假说明", "公司放假加班分钟", "工作日加班分钟", "休息日加班分钟",
            "公司放假抵扣迟到分钟", "当天应出勤分钟",
        ]
        sheet.append(headers)
        for result in sorted(results, key=lambda item: (item.work_date, item.employee_name)):
            cells = _append_row(sheet, _legacy_detail_row(result, self._config))
            if _supplement_punch_issue(result, self._config) is not None:
                _mark_row(cells, RED_FILL, RED_FONT)
        _style_header(sheet)
        sheet.auto_filter.ref = f"A1:{get_column_letter(sheet.max_column)}{sheet.max_row}"

    def _write_employee_summary(
        self, sheet, results: list[OvertimeResult], input_issues: list[InputIssue]
    ) -> None:
        grouped = _group_results(results)
        sheet.append([
            "姓名", "UserId", "统计状态", "阻断输入条数", "待补卡天数", "待审核天数",
            "暂计加班天数", "暂计加班分钟", "暂计加班小时", "迟到次数", "迟到分钟",
            "请假抵扣迟到分钟", "宽限迟到分钟", "扣款迟到分钟",
        ])
        for (name, user_id), items in sorted(grouped.items()):
            _, blocking_issue_count = _person_issue_counts(input_issues, name, user_id)
            supplement_count = sum(_supplement_punch_issue(item, self._config) is not None for item in items)
            review_count = sum(_needs_review(item) for item in items)
            minutes = sum(item.overtime_minutes for item in items)
            row = [
                name, user_id, _summary_status(blocking_issue_count, supplement_count, review_count),
                blocking_issue_count, supplement_count, review_count, sum(item.has_overtime for item in items),
                minutes, round(minutes / 60, 2), sum(item.late_minutes > 0 for item in items),
                sum(item.late_minutes for item in items), sum(item.late_leave_exempt_minutes for item in items),
                sum(item.late_grace_minutes for item in items), sum(item.late_chargeable_minutes for item in items),
            ]
            cells = _append_row(sheet, row)
            if supplement_count:
                _mark_row(cells, RED_FILL, RED_FONT)
        _style_header(sheet)
        sheet.auto_filter.ref = f"A1:{get_column_letter(sheet.max_column)}{sheet.max_row}"

    def _write_daily_summary(
        self, sheet, results: list[OvertimeResult], input_issues: list[InputIssue]
    ) -> None:
        grouped: dict[date, list[OvertimeResult]] = defaultdict(list)
        for result in results:
            grouped[result.work_date].append(result)
        sheet.append([
            "日期", "统计状态", "人工确认条数", "阻断输入条数", "漏卡人数", "待审核人数",
            "加班人数", "加班分钟", "加班小时", "工作日加班分钟", "休息日加班分钟", "节假日加班分钟",
            "请假有效分钟", "扣工资请假分钟", "原始迟到分钟", "请假抵扣迟到分钟", "规则后迟到分钟",
            "宽限迟到分钟", "扣工资迟到分钟",
            "公司放假加班分钟", "公司放假抵扣迟到分钟",
        ])
        for work_date, items in sorted(grouped.items()):
            all_issues = [item for item in input_issues if item.work_date == work_date]
            blocking = sum(item.blocks_summary for item in all_issues)
            supplement = sum(_supplement_punch_issue(item, self._config) is not None for item in items)
            reviews = sum(_needs_review(item) for item in items)
            minutes = sum(item.overtime_minutes for item in items)
            status = _summary_status(blocking, supplement, reviews)
            if status == "有效" and all_issues:
                status = "有效（有待确认项）"
            cells = _append_row(sheet, [
                work_date, status, len(all_issues), blocking, supplement, reviews,
                sum(item.has_overtime for item in items), minutes, round(minutes / 60, 2),
                sum(item.workday_overtime_minutes for item in items),
                sum(item.rest_day_overtime_minutes for item in items),
                sum(item.overtime_minutes for item in items if item.overtime_type == "节假日"),
                sum(item.leave_work_minutes for item in items), sum(item.leave_deductible_minutes for item in items),
                sum(item.late_raw_minutes for item in items), sum(item.late_leave_exempt_minutes for item in items),
                sum(item.late_minutes for item in items), sum(item.late_grace_minutes for item in items),
                sum(item.late_chargeable_minutes for item in items),
                sum(item.company_holiday_overtime_minutes for item in items),
                sum(item.late_company_holiday_exempt_minutes for item in items),
            ])
            if supplement:
                _mark_row(cells, RED_FILL, RED_FONT)
            elif reviews:
                _mark_cells(cells, {1, 4}, AMBER_FILL, AMBER_FONT)
            elif all_issues:
                _mark_cells(cells, {1, 2}, YELLOW_FILL, None)
        _style_header(sheet)
        sheet.auto_filter.ref = f"A1:{get_column_letter(sheet.max_column)}{sheet.max_row}"

    def _write_supplement_punch_notifications(self, sheet, results: list[OvertimeResult]) -> None:
        sheet.append(["姓名", "UserId", "日期", "班次", "需补打卡", "原因", "关联审批单", "建议通知"])
        notification_count = 0
        for result in sorted(results, key=lambda item: (item.work_date, item.employee_name)):
            issue = _supplement_punch_issue(result, self._config)
            if issue is None:
                continue
            missing_type, reason = issue
            cells = _append_row(sheet, [
                result.employee_name, result.user_id, result.work_date, result.shift, missing_type,
                reason, result.approval_text,
                f"请补齐 {result.work_date:%Y-%m-%d} 的{missing_type}；补卡后重新导出考勤并运行本工具。",
            ])
            _mark_row(cells, RED_FILL, RED_FONT)
            notification_count += 1
        if notification_count == 0:
            sheet.append(["未发现待补卡记录"])
        _style_header(sheet)
        sheet.auto_filter.ref = f"A1:{get_column_letter(sheet.max_column)}{sheet.max_row}"

    def _write_approval_reviews(self, sheet, results: list[OvertimeResult]) -> None:
        sheet.append([
            "姓名", "UserId", "日期", "加班类型", "审核类型", "实际首个打卡", "实际最后打卡",
            "实际加班分钟", "请假有效分钟", "请假重叠分钟", "审批加班分钟", "差异(分钟)",
            "容差(分钟)", "关联审批单", "审核原因", "需通知HR",
            "公司放假时段", "公司放假加班分钟", "审批核对范围", "审批核对实际分钟",
        ])
        review_count = 0
        for result in sorted(results, key=lambda item: (item.work_date, item.employee_name)):
            if not _needs_review(result):
                continue
            cells = _append_row(sheet, [
                result.employee_name, result.user_id, result.work_date, result.overtime_type,
                _review_type(result), fmt_dt(result.first_clock_in), fmt_dt(result.last_clock_out),
                result.overtime_minutes, result.leave_work_minutes, result.leave_overlap_minutes,
                result.approval_overtime_minutes, result.approval_difference_minutes,
                self._config.rest_day.approval_difference_tolerance_minutes, result.approval_text,
                _review_reason(result), "是",
                result.company_holiday_period, result.company_holiday_overtime_minutes,
                "仅公司放假时段" if result.overtime_type == "工作日" and result.company_holiday_overtime_minutes else
                ("当日非工作日加班" if result.overtime_type != "工作日" else "工作日加班不要求审批"),
                result.company_holiday_overtime_minutes if result.overtime_type == "工作日" and result.company_holiday_overtime_minutes else result.overtime_minutes,
            ])
            _mark_row(cells, AMBER_FILL, AMBER_FONT)
            review_count += 1
        if review_count == 0:
            sheet.append(["未发现待审核记录"])
        _style_header(sheet)
        sheet.auto_filter.ref = f"A1:{get_column_letter(sheet.max_column)}{sheet.max_row}"

    def _write_anomalies(self, sheet, results: list[OvertimeResult], input_issues: list[InputIssue]) -> None:
        sheet.append([
            "异常类型", "姓名", "UserId", "日期", "班次", "上班打卡结果", "下班打卡结果",
            "原始记录数", "有效原始记录数", "使用每日统计兜底", "关联审批单", "说明",
        ])
        anomaly_count = 0
        for result in sorted(results, key=lambda item: (item.work_date, item.employee_name)):
            for anomaly_type, message in _anomalies_for(result, self._config):
                cells = _append_row(sheet, [
                    anomaly_type, result.employee_name, result.user_id, result.work_date, result.shift,
                    result.start_result, result.end_result, result.raw_punch_count, result.valid_raw_punch_count,
                    "是" if result.used_daily_fallback else "", result.approval_text, message,
                ])
                if "补卡" in anomaly_type or "缺卡" in anomaly_type:
                    _mark_row(cells, RED_FILL, RED_FONT)
                else:
                    _mark_row(cells, AMBER_FILL, AMBER_FONT)
                anomaly_count += 1
        if anomaly_count == 0:
            sheet.append(["未发现异常数据"])
        _style_header(sheet)
        sheet.auto_filter.ref = f"A1:{get_column_letter(sheet.max_column)}{sheet.max_row}"

    def _write_input_issues(self, sheet, input_issues: list[InputIssue]) -> None:
        sheet.append([
            "来源工作表", "源行号", "姓名", "UserId", "考勤日期", "字段", "原始值", "问题",
            "对统计的影响", "建议处理", "问题类型", "阻断自动汇总", "HR选择", "HR处理说明",
        ])
        for issue in input_issues:
            cells = _append_row(sheet, [
                issue.source_sheet, issue.source_row, issue.employee_name, issue.user_id, issue.work_date,
                issue.field_name, issue.raw_value, issue.message, issue.impact, issue.suggested_action,
                issue.issue_type, "是" if issue.blocks_summary else "否", "", "",
            ])
            _mark_row(cells, YELLOW_FILL, None)
        if not input_issues:
            sheet.append(["未发现需要人工确认的输入数据"])
        else:
            choice_validation = DataValidation(
                type="list", formula1='"更正源数据后重跑,确认排除本行,人工核定后重跑"', allow_blank=True
            )
            sheet.add_data_validation(choice_validation)
            choice_validation.add(f"M2:M{len(input_issues) + 1}")
        _style_header(sheet)
        sheet.auto_filter.ref = f"A1:N{sheet.max_row}"
        for row in sheet.iter_rows(min_row=2, min_col=7, max_col=14):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    def _write_config(self, sheet, output_path: Path, input_issues: list[InputIssue]) -> None:
        workday = self._config.workday
        overtime = self._config.overtime
        rest_day = self._config.rest_day
        rows = [
            ("输出文件", str(output_path)),
            ("输入待人工确认总条数", len(input_issues)),
            ("其中阻断自动汇总条数", sum(issue.blocks_summary for issue in input_issues)),
            ("待人工确认处理", "不参与自动猜测；在待人工确认填写 HR 选择后修正源数据并重跑"),
            ("待人工确认 HR选择", "更正源数据后重跑 / 确认排除本行 / 人工核定后重跑"),
            ("输出结构", "总览 → 每位员工整月明细 → 待补卡/待审核/待人工确认/异常/配置"),
            ("弹性上班开始", workday.flex_start_begin.strftime("%H:%M")),
            ("弹性上班结束", workday.flex_start_end.strftime("%H:%M")),
            ("10:00后下班规则", "固定 19:00，不因迟到顺延"),
            ("工作分钟", workday.work_minutes),
            ("午休分钟", workday.lunch_break_minutes),
            ("迟到次数宽限上限(月)", workday.late_grace_occurrences_per_month),
            ("迟到分钟宽限上限(月)", workday.late_grace_minutes_per_month),
            ("是否统计非工作日加班", str(workday.count_non_workday_overtime)),
            ("休息班次关键字", "、".join(workday.rest_shift_keywords)),
            ("缺卡处理", "缺少成对有效打卡时不计入加班；明细整行标红并生成待补卡通知"),
            ("请假处理", "按审批起止时间精确到分钟；有效分钟扣除固定午休；同日加班交 HR 审核"),
            ("不扣工资请假类型", "年假、调休"),
            ("其他请假类型", "事假计入扣工资；病假、婚假、通用请假暂不分类并交 HR 审核；未知审批类型报错"),
            ("外出/出差", "按审批实际区间扣午休计出勤，不计入请假及扣工资请假分钟"),
            ("请假时间依据", "审批区间逐日裁切，扣固定午休，每天最多工作分钟；不按班次09:30裁切，不采用审批尾部小时数"),
            ("请假休息日", "休息日和节假日不计请假分钟"),
            ("输入解析失败", "整体报错并返回非零退出码；列出来源和原始值，不生成本文件的统计结果"),
            (
                "固定午休",
                f"{rest_day.fixed_lunch_start:%H:%M}-"
                f"{(datetime.combine(date.min, rest_day.fixed_lunch_start) + timedelta(minutes=rest_day.fixed_lunch_minutes)):%H:%M}",
            ),
            ("审批与实际差异容差(分钟)", rest_day.approval_difference_tolerance_minutes),
            ("非工作日晚餐时段起点", rest_day.dinner_break_start.strftime("%H:%M")),
            ("非工作日晚餐扣除分钟", rest_day.dinner_break_minutes),
            ("额外节假日日期", "、".join(item.isoformat() for item in rest_day.holiday_dates)),
            ("节假日识别", "仅 holiday_dates 明确配置的法定节假日；其余休息班次为休息日"),
            ("输入月份要求", "同一员工同一自然月须在一份完整文件中；迟到宽限不跨文件分配"),
            ("工作日晚餐扣除分钟", overtime.dinner_break_minutes),
            ("最小加班分钟", overtime.min_overtime_minutes),
            ("取整单位分钟", overtime.rounding_unit_minutes),
            ("取整方式", overtime.rounding_mode),
            ("无效打卡前缀", "、".join(overtime.invalid_punch_prefixes)),
            ("跨午夜归属", "按原始表考勤日期归属前一考勤日"),
            ("工资金额", "当前仅输出扣工资分钟；未配置时薪/日薪，因此不计算金额"),
            ("公司放假适用", "全员；全天免考勤，半天仅对剩余工作时段核算迟到和请假；法定节假日和原休息班次优先"),
            ("公司上午放假", "午休结束到岗，原班次下班；上午工作按休息日加班，正常下班加晚餐后另计工作日加班"),
            ("公司下午放假", "午休开始可下班；之后实际工作按休息日加班扣固定午休及非工作日晚餐，不再叠加工作日加班"),
            ("公司假审批", "放假期间加班需审批；半天公司假仅比较放假时段分钟；与工作日加班分段取整并汇总"),
        ]
        rows.extend((f"公司放假 {item.work_date:%Y-%m-%d}", f"{item.period}；{item.note}") for item in self._config.company_holidays)
        sheet.append(["配置项", "值"])
        for row in rows:
            sheet.append(list(row))
        _style_header(sheet)
        sheet.auto_filter.ref = f"A1:B{sheet.max_row}"


def _group_results(results: list[OvertimeResult]) -> dict[tuple[str, str], list[OvertimeResult]]:
    grouped: dict[tuple[str, str], list[OvertimeResult]] = defaultdict(list)
    for result in results:
        grouped[(result.employee_name, result.user_id)].append(result)
    return grouped


def _period_bounds(results: list[OvertimeResult]) -> tuple[date | None, date | None]:
    if not results:
        return None, None
    return min(item.work_date for item in results), max(item.work_date for item in results)


def _period_text(start: date | None, end: date | None) -> str:
    if start is None or end is None:
        return ""
    return f"{start:%Y-%m}" if start.year == end.year and start.month == end.month else f"{start:%Y-%m-%d}至{end:%Y-%m-%d}"


def _date_range(start: date, end: date) -> list[date]:
    return [start + timedelta(days=offset) for offset in range((end - start).days + 1)]


def _unique_sheet_name(name: str, used: set[str], index: int) -> str:
    safe = INVALID_SHEET_CHARS.sub("_", str(name).strip()).strip() or f"员工{index}"
    safe = safe[:31]
    candidate = safe
    suffix = 2
    while candidate in used:
        tail = f"_{suffix}"
        candidate = f"{safe[:31 - len(tail)]}{tail}"
        suffix += 1
    used.add(candidate)
    return candidate


def _leave_types(groups: object) -> list[str]:
    found: set[str] = set()
    for group in groups:
        for result in group:
            found.update(result.leave_minutes_by_type)
    ordered = [kind for kind in KNOWN_LEAVE_TYPES if kind in found]
    ordered.extend(sorted(found - set(ordered)))
    return ordered


def _person_issue_counts(input_issues: list[InputIssue], employee_name: str, user_id: str) -> tuple[int, int]:
    matched = [
        issue for issue in input_issues
        if issue.employee_name == employee_name
        and (not issue.user_id or not user_id or issue.user_id == user_id)
    ]
    return len(matched), sum(issue.blocks_summary for issue in matched)


def _issues_for_day(input_issues: list[InputIssue], employee_name: str, user_id: str, work_date: date) -> list[InputIssue]:
    return [
        issue for issue in input_issues
        if issue.employee_name == employee_name
        and (not issue.user_id or not user_id or issue.user_id == user_id)
        and issue.work_date == work_date
    ]


def _overview_row(
    name: str,
    user_id: str,
    period_text: str,
    items: list[OvertimeResult],
    all_issue_count: int,
    blocking_issue_count: int,
    config: AppConfig,
    leave_types: list[str],
) -> list[object]:
    minutes = sum(item.overtime_minutes for item in items)
    supplement_count = sum(_supplement_punch_issue(item, config) is not None for item in items)
    review_count = sum(_needs_review(item) for item in items)
    anomaly_dates = {
        item.work_date for item in items
        if _anomalies_for(item, config) or _supplement_punch_issue(item, config) is not None
    }
    type_minutes = {kind: sum(item.leave_minutes_by_type.get(kind, 0) for item in items) for kind in leave_types}
    effective_leave = sum(item.leave_work_minutes for item in items)
    known_minutes = sum(type_minutes.values())
    other_leave = max(effective_leave - known_minutes, 0)
    return [
        name, user_id, period_text,
        _overview_status(blocking_issue_count, supplement_count, review_count, all_issue_count),
        len(items), sum(item.punch_pair_complete for item in items),
        sum(item.has_overtime for item in items), minutes, round(minutes / 60, 2),
        sum(item.workday_overtime_minutes for item in items),
        sum(item.rest_day_overtime_minutes for item in items),
        sum(item.overtime_minutes for item in items if item.overtime_type == "节假日"),
        sum(item.overtime_minutes for item in items if _needs_review(item)), supplement_count,
        sum(item.approved_leave_minutes for item in items), effective_leave,
        *[type_minutes[kind] for kind in leave_types], other_leave,
        sum(item.leave_non_deductible_minutes for item in items),
        sum(item.leave_deductible_minutes for item in items), sum(item.leave_overlap_minutes for item in items),
        sum(item.late_raw_minutes > 0 for item in items), sum(item.late_raw_minutes for item in items),
        sum(item.late_leave_exempt_minutes for item in items), sum(item.late_minutes > 0 for item in items),
        sum(item.late_minutes for item in items), sum(item.late_grace_minutes for item in items),
        sum(item.late_chargeable_minutes > 0 for item in items), sum(item.late_chargeable_minutes for item in items),
        review_count, len(anomaly_dates), all_issue_count, blocking_issue_count,
        "待处理" if all_issue_count else "无",
        sum(bool(item.leave_work_minutes and item.overtime_minutes) for item in items),
    ]


def _overview_friendly_row(
    items: list[OvertimeResult],
    config: AppConfig,
    all_issue_count: int,
    blocking_issue_count: int,
) -> list[object]:
    """Return the compact, reader-facing metrics for one employee."""
    overtime_minutes = sum(item.overtime_minutes for item in items)
    anomaly_dates = {
        item.work_date for item in items
        if _anomalies_for(item, config) or _supplement_punch_issue(item, config) is not None
    }
    review_count = sum(_needs_review(item) for item in items)
    supplement_count = sum(_supplement_punch_issue(item, config) is not None for item in items)
    return [
        overtime_minutes,
        round(overtime_minutes / 60, 2),
        sum(item.workday_overtime_minutes for item in items),
        sum(item.rest_day_overtime_minutes for item in items),
        sum(item.overtime_minutes for item in items if item.overtime_type == "节假日"),
        sum(item.overtime_minutes for item in items if _needs_review(item)),
        sum(item.leave_work_minutes for item in items),
        sum(item.leave_non_deductible_minutes for item in items),
        sum(item.leave_deductible_minutes for item in items),
        sum(item.leave_overlap_minutes for item in items),
        sum(item.late_raw_minutes > 0 for item in items),
        sum(item.late_raw_minutes for item in items),
        sum(item.late_leave_exempt_minutes for item in items),
        sum(item.late_minutes > 0 for item in items),
        sum(item.late_grace_minutes for item in items),
        sum(item.late_chargeable_minutes > 0 for item in items),
        sum(item.late_chargeable_minutes for item in items),
        sum(item.approval_difference_minutes or 0 for item in items),
        review_count,
        supplement_count,
        len(anomaly_dates),
        all_issue_count,
        blocking_issue_count,
        _overview_status(blocking_issue_count, supplement_count, review_count, all_issue_count),
        sum(bool(item.leave_work_minutes and item.overtime_minutes) for item in items),
        sum(item.business_attendance_minutes for item in items),
        sum(item.leave_unclassified_minutes for item in items),
        sum(item.company_holiday_overtime_minutes for item in items),
        sum(item.late_company_holiday_exempt_minutes for item in items),
    ]


def _employee_detail_row(
    result: OvertimeResult | None,
    work_date: date,
    day_issues: list[InputIssue],
    config: AppConfig,
    leave_types: list[str],
) -> list[object]:
    if result is None:
        day_type = _calendar_day_type(work_date, config)
        holiday = config.company_holiday_on(work_date)
        exempt = bool(holiday and holiday.period == "全天")
        return [
            work_date, _weekday_text(work_date), "", day_type, "", "", "", "", "", "", "否", "",
            0, 0, "", "", "", "", 0, 0, 0, 0, 0, "公司全天放假，无需考勤。" if exempt else "未找到该日期的每日统计记录。", "", "", 0, 0,
            *([0] * len(leave_types)), 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, "", "", "", "",
            "" if exempt else "缺少每日统计记录", len(day_issues), "待处理" if day_issues else "无", "正常" if exempt else "缺少每日统计记录",
            "公司放假，免考勤" if exempt else "缺少每日统计记录", "",
        ]

    issue_count = len(day_issues)
    supplement = _supplement_punch_issue(result, config)
    anomaly_types = [kind for kind, _ in _anomalies_for(result, config)]
    if issue_count:
        anomaly_types.append("待人工确认")
    anomaly_status = "、".join(dict.fromkeys(anomaly_types)) or "正常"
    leave_types_text = "、".join(sorted(result.leave_minutes_by_type))
    known_minutes = sum(result.leave_minutes_by_type.get(kind, 0) for kind in leave_types)
    other_leave = max(result.leave_work_minutes - known_minutes, 0)
    cross_midnight = "是" if result.last_clock_out and result.last_clock_out.date() != result.work_date else "否"
    return [
        result.work_date, _weekday_text(result.work_date), result.shift, result.overtime_type,
        result.start_time_text, result.start_result, result.end_time_text, result.end_result,
        fmt_dt(result.first_clock_in), fmt_dt(result.last_clock_out),
        "是" if result.punch_pair_complete else "否", cross_midnight,
        result.raw_punch_count, result.valid_raw_punch_count, "是" if result.used_daily_fallback else "否",
        result.source_row, fmt_time(result.overtime_start), fmt_time(result.overtime_end),
        result.raw_overtime_minutes, result.lunch_break_minutes, result.dinner_break_minutes,
        result.overtime_minutes, result.overtime_hours, _calculation_text(result), result.approval_text,
        leave_types_text, result.approved_leave_minutes, result.leave_work_minutes,
        *[result.leave_minutes_by_type.get(kind, 0) for kind in leave_types], other_leave,
        result.leave_non_deductible_minutes, result.leave_deductible_minutes, result.leave_overlap_minutes,
        result.late_raw_minutes, result.late_leave_exempt_minutes, result.late_minutes,
        result.late_occurrence_number, result.late_grace_minutes, result.late_chargeable_minutes,
        result.late_status, result.approval_overtime_minutes, result.approval_difference_minutes,
        _approval_status(result), supplement[0] if supplement else "", issue_count,
        "待处理" if issue_count else "无", anomaly_status, _calculation_status(result, config),
        "；".join(result.notes),
    ]


def _legacy_detail_row(result: OvertimeResult, config: AppConfig) -> list[object]:
    return [
        result.employee_name, result.user_id, result.work_date, result.shift, result.overtime_type,
        fmt_dt(result.first_clock_in), fmt_dt(result.last_clock_out), fmt_time(result.standard_end),
        result.lunch_break_minutes, result.dinner_break_minutes, fmt_time(result.overtime_start),
        fmt_time(result.overtime_end), result.raw_overtime_minutes, result.overtime_minutes,
        result.overtime_hours, result.approval_text, result.approved_leave_minutes, result.leave_work_minutes,
        "、".join(sorted(result.leave_minutes_by_type)), result.leave_non_deductible_minutes,
        result.leave_deductible_minutes, result.leave_overlap_minutes, result.approval_overtime_minutes,
        result.approval_difference_minutes, _approval_status(result), result.late_raw_minutes,
        result.late_leave_exempt_minutes, result.late_minutes, result.late_occurrence_number,
        result.late_grace_minutes, result.late_chargeable_minutes, result.late_status,
        _calculation_status(result, config), "；".join(result.notes), _calculation_text(result),
        result.company_holiday_period, result.company_holiday_note, result.company_holiday_overtime_minutes,
        result.workday_overtime_minutes, result.rest_day_overtime_minutes,
        result.late_company_holiday_exempt_minutes, result.required_work_minutes,
    ]


def _weekday_text(value: date) -> str:
    return ("一", "二", "三", "四", "五", "六", "日")[value.weekday()]


def _calendar_day_type(value: date, config: AppConfig) -> str:
    if value in config.rest_day.holiday_dates:
        return "节假日"
    holiday = config.company_holiday_on(value)
    if holiday:
        return f"公司放假（{holiday.period}）"
    return "休息日" if value.weekday() >= 5 else "工作日"


def _overview_status(blocking_issue_count: int, supplement_count: int, review_count: int, all_issue_count: int) -> str:
    if blocking_issue_count:
        return "待人工确认，汇总未生效"
    if supplement_count:
        return "待补卡，相关日不计入"
    if review_count:
        return "待审核，暂计"
    if all_issue_count:
        return "有效（有待确认项）"
    return "有效"


def _summary_status(input_issue_count: int, supplement_count: int, review_count: int) -> str:
    if input_issue_count:
        return "待人工确认，汇总未生效"
    if supplement_count:
        return "待补卡，汇总未生效"
    if review_count:
        return "待审核，汇总未生效"
    return "有效"


def _append_row(sheet, values):
    # 直接保留本行单元格供标色，避免 sheet[max_row] 再次扫描整张表。
    cells = tuple(Cell(sheet, value=value) for value in values)
    sheet.append(cells)
    return cells


def _style_header(sheet) -> None:
    header = sheet.cell(1, 1)
    header.fill = HEADER_FILL
    header.font = Font(color="FFFFFF", bold=True)
    header.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    header.border = Border(bottom=Side(style="medium", color="17365D"))
    for cell in sheet[1]:
        cell._style = copy(header._style)
    sheet.row_dimensions[1].height = 34


def _mark_row(row, fill: PatternFill, font: Font | None) -> None:
    styles = {}
    for cell in row:
        key = tuple(cell._style) if cell.has_style else None
        if key not in styles:
            cell.fill = fill
            if font is not None:
                cell.font = font
            styles[key] = copy(cell._style)
        else:
            cell._style = copy(styles[key])


def _mark_cells(row, indexes: set[int], fill: PatternFill, font: Font | None) -> None:
    for index in indexes:
        if 0 <= index < len(row):
            row[index].fill = fill
            if font is not None:
                row[index].font = font


def _finish_sheet(sheet) -> None:
    max_row, max_column = sheet.max_row, sheet.max_column
    headers = next(sheet.iter_rows(min_row=1, max_row=1, max_col=max_column))
    number_formats = {
        cell.column: "@" if cell.value == "UserId" else "yyyy-mm-dd"
        for cell in headers if cell.value in {"UserId", "日期", "考勤日期"}
    }
    sheet.sheet_view.showGridLines = False
    if sheet.title == "总览":
        sheet.freeze_panes = "D2"
    elif sheet.title in {"待人工确认", "配置"}:
        sheet.freeze_panes = "A2"
    else:
        sheet.freeze_panes = "E2" if max_column > 25 else "A2"

    # 在复用最终样式前一并设置日期/文本格式，避免共享样式被后续修改。
    styles = {}
    for row in sheet.iter_rows(min_row=2, max_row=max_row, max_col=max_column):
        for cell in row:
            wrap_text = isinstance(cell.value, str) and len(cell.value) > 24
            number_format = number_formats.get(cell.column)
            key = (tuple(cell._style) if cell.has_style else None, wrap_text, number_format)
            if key not in styles:
                cell.border = GRID_BORDER
                if wrap_text:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                if number_format:
                    cell.number_format = number_format
                styles[key] = cell._style
            else:
                cell._style = styles[key]

    header_by_col = {cell.column: str(cell.value) for cell in headers}
    for column_cells in sheet.columns:
        column_letter = get_column_letter(column_cells[0].column)
        header = header_by_col.get(column_cells[0].column, "")
        max_length = max((len(str(cell.value)) for cell in column_cells if cell.value is not None), default=8)
        if header in {"关联审批单", "加班计算说明", "计算说明", "备注", "问题", "对统计的影响", "建议处理", "审核原因", "HR处理说明"}:
            width = 34
        elif "状态" in header or header in {"统计期间", "异常状态", "请假类型", "漏卡", "人工确认状态"}:
            width = 20
        elif header == "UserId":
            width = 24
        elif "分钟" in header or "次数" in header or "天数" in header or "记录数" in header:
            width = 14
        else:
            width = min(max(max_length + 2, 10), 22)
        sheet.column_dimensions[column_letter].width = min(max(width, 10), 42)

    if sheet.title == "待人工确认":
        sheet.column_dimensions["M"].width = 24
        sheet.column_dimensions["N"].width = 32
        for row in range(2, max_row + 1):
            sheet.row_dimensions[row].height = 48


def _calculation_text(result: OvertimeResult) -> str:
    leave_text = _leave_calculation_text(result)
    if result.business_attendance_minutes:
        leave_text += f"；外出/出差按审批区间扣午休计出勤 {result.business_attendance_minutes} 分钟"
    if result.company_holiday_period == "全天" and not result.has_attendance_signal:
        return "公司全天放假，免考勤，不计迟到或请假；无实际加班。"
    if result.company_holiday_period and result.overtime_type == "工作日":
        return (
            f"公司{result.company_holiday_period}放假；当天应出勤 {result.required_work_minutes} 分钟，"
            f"标准下班 {fmt_time(result.standard_end)}；公司放假时段按休息日计 {result.company_holiday_overtime_minutes} 分钟"
            f"（{fmt_time(result.company_holiday_overtime_start)}至{fmt_time(result.company_holiday_overtime_end)}），"
            f"工作日加班另计 {result.workday_overtime_minutes} 分钟；合计 {result.overtime_minutes} 分钟；"
            f"公司放假抵扣迟到 {result.late_company_holiday_exempt_minutes} 分钟，规则后迟到 {result.late_minutes} 分钟{leave_text}。"
        )
    if result.overtime_type != "工作日" and result.first_clock_in and result.last_clock_out:
        deduction = f"固定午休扣除 {result.lunch_break_minutes} 分钟"
        if result.dinner_break_minutes:
            deduction += f"，晚餐扣除 {result.dinner_break_minutes} 分钟"
        return f"{result.overtime_type}首个有效打卡 {fmt_time(result.first_clock_in)}，最后有效打卡 {fmt_time(result.last_clock_out)}；{deduction}；计入 {result.overtime_minutes} 分钟{leave_text}。"
    if result.first_clock_in is None or result.last_clock_out is None or result.standard_end is None:
        notes = "；".join(result.notes) or "缺少有效上下班打卡"
        return f"未计算：{notes}{leave_text}。"
    late_text = ""
    if result.late_minutes:
        late_text = f"；规则后迟到 {result.late_minutes} 分钟，宽限 {result.late_grace_minutes} 分钟，扣工资 {result.late_chargeable_minutes} 分钟"
    return f"首个有效打卡 {fmt_time(result.first_clock_in)}，弹性下班 {fmt_time(result.standard_end)}；晚餐扣除 {result.dinner_break_minutes} 分钟后 {fmt_time(result.overtime_start)} 开始算加班；最后有效打卡 {fmt_time(result.last_clock_out)}；计入 {result.overtime_minutes} 分钟{late_text}{leave_text}。"


def _leave_calculation_text(result: OvertimeResult) -> str:
    if not result.approved_leave_minutes:
        return ""
    text = f"；请假审批区间覆盖 {result.approved_leave_minutes} 分钟，按当天区间扣午休并限制日上限后有效 {result.leave_work_minutes} 分钟"
    if result.late_leave_exempt_minutes:
        text += f"，抵扣迟到 {result.late_leave_exempt_minutes} 分钟"
    if result.leave_overlap_minutes:
        text += f"，重叠 {result.leave_overlap_minutes} 分钟待 HR 确认"
    return text


def _anomalies_for(result: OvertimeResult, config: AppConfig) -> list[tuple[str, str]]:
    anomalies: list[tuple[str, str]] = []
    start_missing = not result.start_time_text and bool(result.end_time_text)
    end_missing = bool(result.start_time_text) and not result.end_time_text
    approved_non_punch = _has_approved_non_punch_reason(result, config)
    supplement_issue = _supplement_punch_issue(result, config)
    if supplement_issue is not None:
        anomalies.append(("待补卡", supplement_issue[1]))
    company_pair_confirmed = bool(result.company_holiday_period and result.punch_pair_complete)
    if start_missing and not approved_non_punch and not company_pair_confirmed:
        anomalies.append(("有下班无上班", "每日统计中缺少上班打卡时间，但存在下班打卡时间。"))
    if end_missing and not approved_non_punch and not company_pair_confirmed:
        anomalies.append(("有上班无下班", "每日统计中存在上班打卡时间，但缺少下班打卡时间。"))
    if ("缺卡" in result.start_result or "缺卡" in result.end_result) and result.raw_punch_count > 0 and not company_pair_confirmed:
        anomalies.append(("缺卡但有原始记录", "每日统计显示缺卡，但原始记录中存在该员工当天打卡流水。"))
    if result.raw_punch_count > 0 and result.valid_raw_punch_count == 0:
        anomalies.append(("原始记录全部无效", "原始记录存在打卡流水，但都被无效打卡规则排除。"))
    if result.approval_review_reason:
        anomalies.append(("审批与实际打卡待审核", result.approval_review_reason))
    if result.leave_review_reason:
        anomalies.append(("请假与加班待审核", result.leave_review_reason))
    if _is_non_workday(result, config) and result.has_attendance_signal and not config.workday.count_non_workday_overtime:
        anomalies.append(("非工作日有打卡但未统计", "当天为非工作日，配置为不统计非工作日加班。"))
    has_any_attendance_signal = result.has_attendance_signal
    should_expect_punch_pair = bool(result.shift.strip()) and not _is_non_workday(result, config)
    if any(note == "有效上下班打卡不足" for note in result.notes) and (has_any_attendance_signal or should_expect_punch_pair) and not approved_non_punch and supplement_issue is None:
        anomalies.append(("有效上下班打卡不足", "无法从原始记录或每日统计中得到成对的有效上下班打卡。"))
    if result.leave_overlap_minutes:
        anomalies.append(("请假类型重叠", f"不同请假类型重叠 {result.leave_overlap_minutes} 分钟：{result.leave_overlap_types}。"))
    return anomalies


def _has_approved_non_punch_reason(result: OvertimeResult, config: AppConfig) -> bool:
    required = result.required_work_minutes if result.required_work_minutes is not None else config.workday.work_minutes
    return required > 0 and result.approved_non_punch_minutes >= required


def _supplement_punch_issue(result: OvertimeResult, config: AppConfig) -> tuple[str, str] | None:
    if result.punch_pair_complete or _has_approved_non_punch_reason(result, config):
        return None
    if not _requires_punch_pair(result, config):
        return None
    if not result.start_time_text and result.end_time_text:
        missing_type = "上班打卡"
    elif result.start_time_text and not result.end_time_text:
        missing_type = "下班打卡"
    else:
        missing_type = "上班和下班打卡"
    return missing_type, "缺少成对有效打卡，当天加班不计入，补卡后需重新统计。"


def _requires_punch_pair(result: OvertimeResult, config: AppConfig) -> bool:
    if not result.shift.strip():
        return False
    if not _is_non_workday(result, config):
        return True
    return result.has_attendance_signal


def _calculation_status(result: OvertimeResult, config: AppConfig) -> str:
    if _supplement_punch_issue(result, config) is not None:
        return "待补卡，不计入"
    if _needs_review(result):
        return "待审核，暂计"
    if _is_non_workday(result, config) and not config.workday.count_non_workday_overtime:
        return "休息日未统计"
    if result.company_holiday_period == "全天" and not result.has_attendance_signal:
        return "公司放假，免考勤"
    if result.first_clock_in is None or result.last_clock_out is None:
        return "未计算"
    return "有效"


def _approval_status(result: OvertimeResult) -> str:
    if _needs_review(result):
        return "待审核"
    if result.approval_overtime_minutes is not None:
        return "已核对"
    return ""


def _needs_review(result: OvertimeResult) -> bool:
    return bool(result.approval_review_reason or result.leave_review_reason)


def _review_type(result: OvertimeResult) -> str:
    types: list[str] = []
    if result.approval_review_reason:
        types.append("加班审批差异")
    if result.leave_review_reason:
        types.append("请假与加班/类型重叠")
    return "、".join(types)


def _review_reason(result: OvertimeResult) -> str:
    return "；".join(reason for reason in (result.approval_review_reason, result.leave_review_reason) if reason)


def _input_issue_counts_by_date(input_issues: list[InputIssue]) -> Counter:
    counts: Counter = Counter()
    for issue in input_issues:
        if issue.work_date is not None:
            counts[issue.work_date] += 1
    return counts


def _is_rest_day(result: OvertimeResult, config: AppConfig) -> bool:
    return any(keyword in result.shift for keyword in config.workday.rest_shift_keywords)


def _is_non_workday(result: OvertimeResult, config: AppConfig) -> bool:
    return result.overtime_type != "工作日" or _is_rest_day(result, config) or result.work_date in config.rest_day.holiday_dates
