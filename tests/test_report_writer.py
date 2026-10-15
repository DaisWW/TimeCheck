from datetime import date
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook

from attendance_analyzer.config import AppConfig, RestDayConfig
from attendance_analyzer.models import InputIssue, OvertimeResult
from attendance_analyzer.report_writer import OvertimeReportWriter, _supplement_punch_issue


def _result(work_date: date, **overrides) -> OvertimeResult:
    values = {
        "employee_name": "Alice",
        "user_id": "u1",
        "work_date": work_date,
        "shift": "A",
        "first_clock_in": None,
        "last_clock_out": None,
        "standard_end": None,
        "dinner_break_minutes": 30,
        "overtime_start": None,
        "overtime_end": None,
        "raw_overtime_minutes": 0,
        "overtime_minutes": 0,
        "approval_text": "",
    }
    values.update(overrides)
    return OvertimeResult(**values)


class OvertimeReportWriterTest(unittest.TestCase):
    def test_configured_holiday_without_attendance_does_not_require_punches(self):
        config = AppConfig(rest_day=RestDayConfig(holiday_dates=(date(2026, 10, 1),)))
        result = _result(date(2026, 10, 1), overtime_type="节假日", notes=["有效上下班打卡不足"])
        self.assertIsNone(_supplement_punch_issue(result, config))

    def test_attendance_and_pending_leave_categories_remain_visible(self):
        result = _result(
            date(2026, 8, 3), business_attendance_minutes=150,
            leave_work_minutes=60, leave_unclassified_minutes=60,
            leave_review_reason="病假工资分类待确认", punch_pair_complete=True,
        )
        issue = InputIssue(
            "每日统计", 5, "Alice", "u1", date(2026, 8, 3), "关联的审批单",
            "补卡申请", "补卡待核对", "待确认", "核对后重跑", blocks_summary=False,
        )
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "result.xlsx"
            OvertimeReportWriter(AppConfig()).write(path, [result], [issue])
            workbook = load_workbook(path, data_only=True)
            overview = {cell.value: workbook["总览"].cell(2, cell.column).value for cell in workbook["总览"][1]}
            self.assertEqual(overview["外出/出差出勤时间（分钟）"], 150)
            self.assertEqual(overview["工资分类待确认请假（分钟）"], 60)
            daily = workbook["每日汇总"]
            self.assertIn("待审核", daily["B2"].value)
            workbook.close()

    def test_sick_leave_does_not_create_missing_punch_notification(self):
        result = OvertimeResult(
            employee_name="Alice",
            user_id="u1",
            work_date=date(2026, 7, 17),
            shift="A 09:30-18:30",
            first_clock_in=None,
            last_clock_out=None,
            standard_end=None,
            dinner_break_minutes=30,
            overtime_start=None,
            overtime_end=None,
            raw_overtime_minutes=0,
            overtime_minutes=0,
            approval_text="病假07-17 08:30到07-17 18:30 7.5小时",
            start_result="请假",
            end_result="请假",
            punch_pair_complete=False,
            approved_non_punch_minutes=450,
            notes=["有效上下班打卡不足"],
        )

        self.assertIsNone(_supplement_punch_issue(result, AppConfig()))

    def test_missing_punch_creates_notification_and_invalidates_employee_summary(self):
        result = OvertimeResult(
            employee_name="Alice",
            user_id="u1",
            work_date=date(2026, 5, 6),
            shift="A 09:30-18:30",
            first_clock_in=None,
            last_clock_out=None,
            standard_end=None,
            dinner_break_minutes=30,
            overtime_start=None,
            overtime_end=None,
            raw_overtime_minutes=0,
            overtime_minutes=0,
            approval_text="",
            start_time_text="09:30",
            raw_punch_count=1,
            valid_raw_punch_count=1,
            punch_pair_complete=False,
            notes=["有效上下班打卡不足"],
        )

        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "result.xlsx"
            OvertimeReportWriter(AppConfig()).write(output_path, [result])
            workbook = load_workbook(output_path, data_only=True)

        notification = workbook["待补卡通知"]
        summary = workbook["员工汇总"]
        self.assertEqual(notification["A2"].value, "Alice")
        self.assertEqual(notification["E2"].value, "下班打卡")
        self.assertEqual(summary["C2"].value, "待补卡，汇总未生效")
        self.assertEqual(summary["E2"].value, 1)

    def test_report_separates_late_minutes_and_removes_overtime_cap_columns(self):
        result = OvertimeResult(
            employee_name="Alice",
            user_id="u1",
            work_date=date(2026, 7, 1),
            shift="A",
            first_clock_in=None,
            last_clock_out=None,
            standard_end=None,
            dinner_break_minutes=30,
            overtime_start=None,
            overtime_end=None,
            raw_overtime_minutes=0,
            overtime_minutes=0,
            approval_text="",
            late_minutes=45,
            late_occurrence_number=1,
            late_grace_minutes=30,
            late_chargeable_minutes=15,
            late_status="第1次迟到：宽限 30 分钟，扣款 15 分钟",
        )

        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "result.xlsx"
            OvertimeReportWriter(AppConfig()).write(output_path, [result])
            workbook = load_workbook(output_path, data_only=True)

        detail_header = [cell.value for cell in workbook["加班明细"][1]]
        review_header = [cell.value for cell in workbook["待审核清单"][1]]
        self.assertNotIn("是否封顶", detail_header)
        self.assertIn("扣款迟到分钟", detail_header)
        self.assertIn("需通知HR", review_header)

    def test_input_issue_is_written_to_manual_review_and_invalidates_summary(self):
        result = OvertimeResult(
            employee_name="Alice",
            user_id="u1",
            work_date=date(2026, 7, 1),
            shift="A",
            first_clock_in=None,
            last_clock_out=None,
            standard_end=None,
            dinner_break_minutes=30,
            overtime_start=None,
            overtime_end=None,
            raw_overtime_minutes=0,
            overtime_minutes=0,
            approval_text="",
        )
        input_issue = InputIssue(
            source_sheet="原始记录",
            source_row=12,
            employee_name="Alice",
            user_id="u1",
            work_date=date(2026, 7, 1),
            field_name="打卡时间",
            raw_value="25:00",
            message="原始打卡时间为空或无法解析。",
            impact="该打卡不参与统计。",
            suggested_action="修正打卡时间后重新导出并运行。",
        )

        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "result.xlsx"
            OvertimeReportWriter(AppConfig()).write(output_path, [result], [input_issue])
            workbook = load_workbook(output_path, data_only=True)

        manual_sheet = workbook["待人工确认"]
        summary = workbook["员工汇总"]
        self.assertEqual(manual_sheet["A2"].value, "原始记录")
        self.assertEqual(manual_sheet["G2"].value, "25:00")
        self.assertEqual(manual_sheet.auto_filter.ref, "A1:N2")
        self.assertTrue(manual_sheet["G2"].alignment.wrap_text)
        self.assertEqual(len(manual_sheet.data_validations.dataValidation), 1)
        validation = manual_sheet.data_validations.dataValidation[0]
        self.assertIn("更正源数据后重跑", validation.formula1)
        self.assertEqual(str(validation.sqref), "M2")
        self.assertEqual(summary["C2"].value, "待人工确认，汇总未生效")
        self.assertEqual(summary["D2"].value, 1)

    def test_employee_detail_has_full_month_and_marks_missing_punch_red(self):
        results = [
            _result(date(2026, 5, 1)),
            _result(
                date(2026, 5, 31),
                start_time_text="09:00",
                raw_punch_count=1,
                valid_raw_punch_count=1,
                punch_pair_complete=False,
                notes=["有效上下班打卡不足"],
            ),
        ]

        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "result.xlsx"
            OvertimeReportWriter(AppConfig()).write(output_path, results)
            workbook = load_workbook(output_path, data_only=True)

        employee = workbook["Alice"]
        headers = [cell.value for cell in employee[1]]
        leak_column = headers.index("漏卡") + 1
        self.assertEqual(employee.max_row, 32)
        self.assertEqual(employee.cell(2, 1).value.date(), date(2026, 5, 1))
        self.assertEqual(employee.cell(32, 1).value.date(), date(2026, 5, 31))
        self.assertEqual(employee.cell(32, leak_column).value, "下班打卡")
        self.assertEqual(employee.cell(32, 1).fill.fill_type, "solid")
        self.assertTrue(employee.cell(32, 1).fill.fgColor.rgb.endswith("FFC7CE"))

    def test_overview_exposes_each_late_calculation_stage(self):
        result = _result(
            date(2026, 7, 1),
            late_raw_minutes=45,
            late_leave_exempt_minutes=5,
            late_minutes=40,
            late_occurrence_number=1,
            late_grace_minutes=30,
            late_chargeable_minutes=10,
        )

        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "result.xlsx"
            OvertimeReportWriter(AppConfig()).write(output_path, [result])
            workbook = load_workbook(output_path, data_only=True)

        overview = workbook["总览"]
        values = {cell.value: overview.cell(2, index).value for index, cell in enumerate(overview[1], 1)}
        self.assertEqual(values["原始迟到次数"], 1)
        self.assertEqual(values["原始迟到分钟"], 45)
        self.assertEqual(values["请假抵扣迟到分钟"], 5)
        self.assertEqual(values["规则后迟到次数"], 1)
        self.assertEqual(values["规则后迟到分钟"], 40)
        self.assertEqual(values["宽限迟到分钟"], 30)
        self.assertEqual(values["扣工资迟到次数"], 1)
        self.assertEqual(values["扣工资迟到分钟"], 10)
        self.assertEqual(values["迟到次数（原始）"], 1)
        self.assertEqual(values["迟到时间（分钟）"], 45)
        self.assertEqual(values["迟到计入时间（分钟）"], 10)
        self.assertIn("加班时间（分钟）", values)
        self.assertIn("请假时间（调休、年假）（分钟）", values)
        self.assertIn("请假时间（其他）（分钟）", values)


if __name__ == "__main__":
    unittest.main()
