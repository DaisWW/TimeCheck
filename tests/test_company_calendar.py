from datetime import date, datetime, time
import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from attendance_analyzer import cli
from attendance_analyzer.company_calendar import CompanyHoliday
from attendance_analyzer.config import AppConfig, RestDayConfig, WorkdayConfig
from attendance_analyzer.models import DailyAttendance, PunchRecord
from attendance_analyzer.report_writer import OvertimeReportWriter, _supplement_punch_issue
from attendance_analyzer.rules import OvertimeRule


DAY = date(2026, 10, 19)


def calendar_config(period, **overrides):
    return AppConfig(company_holidays=(CompanyHoliday(DAY, period, "公司活动"),), **overrides)


def calculate(period, start=None, end=None, approval="", **config_overrides):
    config = calendar_config(period, **config_overrides)
    day = DailyAttendance("Alice", "u1", DAY, "A 09:30-18:30", approval_text=approval)
    punches = [
        PunchRecord("Alice", "u1", datetime.combine(DAY, time.fromisoformat(value)), "正常", index)
        for index, value in enumerate((start, end), 1) if value
    ]
    return config, OvertimeRule(config).calculate(day, punches)


class CompanyCalendarInputTest(unittest.TestCase):
    def test_single_json_loads_and_sorts_company_holidays(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "config.json"
            data = AppConfig().to_dict()
            data["company_holidays"] = [
                {"date": "2026-10-20", "period": "下午", "note": "活动"},
                {"date": "2026-10-19", "period": "上午", "note": ""},
            ]
            path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
            self.assertEqual(AppConfig.load(path).company_holidays, (
                CompanyHoliday(DAY, "上午"), CompanyHoliday(date(2026, 10, 20), "下午", "活动"),
            ))

    def test_empty_calendar_does_not_load_examples_or_legacy_companion_files(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            config_path = root / "config.json"
            AppConfig().save(config_path)
            for name in ("公司放假.csv", "公司放假.example.csv"):
                (root / name).write_text("日期,时段,备注\n2026-10-19,全天,示例\n", encoding="utf-8")
            self.assertEqual(AppConfig.load(config_path).company_holidays, ())
            self.assertEqual(AppConfig.load(None).company_holidays, ())

    def test_bad_dates_periods_duplicates_types_and_fields_report_source(self):
        valid = {"date": "2026-10-19", "period": "上午", "note": ""}
        cases = (
            ([valid | {"date": "2026-02-30"}], "非法日期"),
            ([valid | {"date": "26-10-19"}], "YYYY-MM-DD"),
            ([valid | {"date": "2026-10-19备注"}], "YYYY-MM-DD"),
            ([valid | {"period": "半天"}], "时段"),
            ([valid, valid | {"period": "下午"}], "重复或冲突"),
            ([{"date": "2026-10-19", "period": "全天"}], "必须包含"),
            ([valid | {"note": 3}], "备注"),
            ([valid | {"unexpected": True}], "必须包含"),
            ({}, "列表"),
        )
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "config.json"
            for entries, reason in cases:
                with self.subTest(entries=entries):
                    data = AppConfig().to_dict()
                    data["company_holidays"] = entries
                    path.write_text(json.dumps(data), encoding="utf-8")
                    with self.assertRaises(ValueError) as raised:
                        AppConfig.load(path)
                    for fragment in (str(path), "company_holidays", "原值", reason):
                        self.assertIn(fragment, str(raised.exception))

    def test_invalid_calendar_stops_cli_before_reading_input(self):
        with tempfile.TemporaryDirectory() as directory:
            config = Path(directory) / "config.json"
            data = AppConfig().to_dict()
            data["company_holidays"] = [{"date": "2026-10-19", "period": "夜里", "note": ""}]
            config.write_text(json.dumps(data), encoding="utf-8")
            with patch("attendance_analyzer.cli.AttendanceAnalyzer") as analyzer, patch("sys.stderr", new_callable=io.StringIO) as stderr:
                code = cli.main(["-c", str(config)])
            self.assertEqual(code, 2)
            analyzer.assert_not_called()
            self.assertIn("company_holidays[0]", stderr.getvalue())

    def test_cli_applies_hr_json_to_workbook_and_monthly_grace(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            config_path = root / "config.json"
            calendar_config("上午").save(config_path)
            input_path = root / "考勤.xlsx"
            output_path = root / "结果.xlsx"
            workbook = Workbook()
            daily = workbook.active
            daily.title = "每日统计"
            daily.append([])
            daily.append([])
            daily.append(["姓名", "UserId", "日期", "班次", "上班1打卡时间", "上班1打卡结果", "下班1打卡时间", "下班1打卡结果", "关联的审批单"])
            daily.append([])
            daily.append(["Alice", "u1", "2026-10-19", "A 09:30-18:30", "13:45", "迟到", "18:30", "正常", ""])
            raw = workbook.create_sheet("原始记录")
            raw.append([])
            raw.append([])
            raw.append(["姓名", "UserId", "考勤日期", "打卡时间", "打卡结果"])
            for value in ("13:45", "18:30"):
                raw.append(["Alice", "u1", "2026-10-19", f"2026-10-19 {value}", "正常"])
            workbook.save(input_path)
            workbook.close()
            with patch("sys.stdout", new_callable=io.StringIO) as stdout:
                code = cli.main([str(input_path), "-c", str(config_path), "-o", str(output_path)])
            self.assertEqual(code, 0)
            self.assertIn(str(config_path), stdout.getvalue())
            result = load_workbook(output_path, data_only=True)
            try:
                row = {cell.value: result["Alice"].cell(2, cell.column).value for cell in result["Alice"][1]}
                self.assertEqual(row["公司放假时段"], "上午")
                self.assertEqual(row["公司放假抵扣迟到分钟"], 120)
                self.assertEqual(row["规则后迟到分钟"], 15)
                self.assertEqual(row["宽限迟到分钟"], 15)
                self.assertEqual(row["扣工资迟到分钟"], 0)
            finally:
                result.close()


class CompanyHolidayRuleTest(unittest.TestCase):
    def test_full_day_exempts_workday_leave_lateness_and_punches(self):
        config, result = calculate("全天", approval="事假10-19 09:00到10-19 18:30")
        self.assertEqual(result.overtime_type, "休息日")
        self.assertEqual(result.leave_work_minutes, 0)
        self.assertEqual(result.required_work_minutes, 0)
        self.assertEqual(result.late_minutes, 0)
        self.assertIsNone(_supplement_punch_issue(result, config))

    def test_full_day_work_uses_rest_day_approval_and_lunch(self):
        _, result = calculate("全天", "09:30", "18:30", "加班10-19 09:30到10-19 18:30")
        self.assertEqual(result.overtime_minutes, 450)
        self.assertEqual(result.rest_day_overtime_minutes, 450)
        self.assertEqual(result.workday_overtime_minutes, 0)
        self.assertEqual(result.company_holiday_overtime_minutes, 450)
        self.assertEqual(result.approval_difference_minutes, 0)
        self.assertEqual(result.approval_review_reason, "")

    def test_morning_holiday_keeps_shift_end_and_counts_1345_as_15_minutes_late(self):
        _, result = calculate("上午", "13:45", "18:30")
        self.assertEqual(result.standard_end, datetime(2026, 10, 19, 18, 30))
        self.assertEqual(result.required_work_minutes, 300)
        self.assertEqual(result.late_raw_minutes, 135)
        self.assertEqual(result.late_company_holiday_exempt_minutes, 120)
        self.assertEqual(result.late_leave_exempt_minutes, 0)
        self.assertEqual(result.late_minutes, 15)
        self.assertEqual(result.overtime_minutes, 0)
        self.assertEqual(result.approval_review_reason, "")

    def test_morning_holiday_and_approved_leave_do_not_double_exempt_lateness(self):
        _, result = calculate("上午", "13:45", "18:30", "事假10-19 10:00到10-19 13:40")
        self.assertEqual(result.late_company_holiday_exempt_minutes, 120)
        self.assertEqual(result.late_leave_exempt_minutes, 10)
        self.assertEqual(result.late_minutes, 5)
        self.assertEqual(result.leave_work_minutes, 10)

    def test_morning_and_evening_work_are_split_and_only_morning_needs_approval(self):
        _, result = calculate("上午", "09:30", "22:00", "加班10-19 09:30到10-19 22:00")
        self.assertEqual(result.company_holiday_overtime_minutes, 150)
        self.assertEqual(result.workday_overtime_minutes, 180)
        self.assertEqual(result.rest_day_overtime_minutes, 150)
        self.assertEqual(result.overtime_minutes, 330)
        self.assertEqual(result.approval_overtime_minutes, 150)
        self.assertEqual(result.approval_difference_minutes, 0)
        self.assertEqual(result.approval_review_reason, "")

    def test_missing_or_outside_company_time_approval_is_flagged(self):
        for approval in ("", "加班10-19 19:00到10-19 22:00"):
            with self.subTest(approval=approval):
                _, result = calculate("上午", "09:30", "22:00", approval)
                self.assertEqual(result.overtime_minutes, 330)
                self.assertTrue(result.approval_review_reason)

    def test_afternoon_holiday_counts_5_hours_and_continues_into_evening_once(self):
        for end, expected in (("18:30", 300), ("22:00", 510)):
            with self.subTest(end=end):
                _, result = calculate("下午", "09:30", end, f"加班10-19 13:30到10-19 {end}")
                self.assertEqual(result.standard_end, datetime(2026, 10, 19, 12))
                self.assertEqual(result.overtime_start, datetime(2026, 10, 19, 13, 30))
                self.assertEqual(result.company_holiday_overtime_minutes, expected)
                self.assertEqual(result.rest_day_overtime_minutes, expected)
                self.assertEqual(result.workday_overtime_minutes, 0)
                self.assertEqual(result.overtime_minutes, expected)
                self.assertEqual(result.dinner_break_minutes, 0)
                self.assertEqual(result.approval_difference_minutes, 0)

    def test_afternoon_holiday_can_leave_at_noon_without_overtime_or_approval(self):
        _, result = calculate("下午", "09:30", "12:00")
        self.assertEqual(result.overtime_minutes, 0)
        self.assertEqual(result.required_work_minutes, 180)
        self.assertEqual(result.late_minutes, 0)
        self.assertEqual(result.approval_review_reason, "")

    def test_half_day_leave_counts_only_remaining_work_time(self):
        for period, approval, expected in (
            ("上午", "事假10-19 09:00到10-19 18:30", 300),
            ("下午", "事假10-19 09:00到10-19 18:00", 180),
            ("上午", "外出10-19 09:00到10-19 18:30", 300),
        ):
            with self.subTest(period=period, approval=approval):
                config, result = calculate(period, approval=approval)
                if approval.startswith("外出"):
                    self.assertEqual(result.business_attendance_minutes, expected)
                    self.assertEqual(result.leave_work_minutes, 0)
                else:
                    self.assertEqual(result.leave_work_minutes, expected)
                    self.assertEqual(result.leave_deductible_minutes, expected)
                    self.assertIsNone(_supplement_punch_issue(result, config))

    def test_remaining_half_day_still_requires_punches(self):
        config, result = calculate("上午", "13:30")
        self.assertIsNotNone(_supplement_punch_issue(result, config))

    def test_partial_approval_does_not_exempt_remaining_work_without_punches(self):
        config = calendar_config("上午")
        for kind in ("事假", "外出"):
            with self.subTest(kind=kind):
                day = DailyAttendance(
                    "Alice", "u1", DAY, "A 09:30-18:30",
                    start_result=kind, end_result=kind,
                    approval_text=f"{kind}10-19 13:30到10-19 14:30",
                )
                result = OvertimeRule(config).calculate(day, [])
                self.assertIsNotNone(_supplement_punch_issue(result, config))

    def test_full_remaining_work_approval_union_exempts_punches_without_double_counting(self):
        config = calendar_config("上午")
        for approval, exempt in (
            ("外出10-19 13:30到10-19 18:30", True),
            ("事假10-19 13:30到10-19 15:30，外出10-19 15:30到10-19 18:30", True),
            ("事假10-19 13:30到10-19 16:30，外出10-19 13:30到10-19 16:30", False),
        ):
            with self.subTest(approval=approval):
                day = DailyAttendance("Alice", "u1", DAY, "A 09:30-18:30", approval_text=approval)
                result = OvertimeRule(config).calculate(day, [])
                self.assertEqual(_supplement_punch_issue(result, config) is None, exempt)

    def test_afternoon_overtime_crossing_midnight_is_counted_once(self):
        config = calendar_config("下午")
        day = DailyAttendance("Alice", "u1", DAY, "A 09:30-18:30", approval_text="加班10-19 13:30到10-20 01:00")
        punches = [
            PunchRecord("Alice", "u1", datetime(2026, 10, 19, 9, 30), "正常", 1),
            PunchRecord("Alice", "u1", datetime(2026, 10, 20, 1), "正常", 2),
        ]
        result = OvertimeRule(config).calculate(day, punches)
        self.assertEqual(result.overtime_minutes, 690)
        self.assertEqual(result.workday_overtime_minutes, 0)
        self.assertEqual(result.approval_difference_minutes, 0)

    def test_monthly_grace_only_consumes_lateness_after_company_exemption(self):
        from attendance_analyzer.analyzer import AttendanceAnalyzer

        config, result = calculate("上午", "13:45", "18:30")
        allocated = AttendanceAnalyzer(config)._allocate_monthly_late_grace([result])[0]
        self.assertEqual(allocated.late_occurrence_number, 1)
        self.assertEqual(allocated.late_grace_minutes, 15)
        self.assertEqual(allocated.late_chargeable_minutes, 0)

    def test_statutory_holiday_and_existing_rest_shift_keep_whole_day_rules(self):
        _, result = calculate("下午", "09:30", "18:30", rest_day=RestDayConfig(holiday_dates=(DAY,)))
        self.assertEqual(result.overtime_type, "节假日")
        self.assertEqual(result.overtime_minutes, 450)
        config = calendar_config("上午")
        day = DailyAttendance("Alice", "u1", DAY, "休息")
        punches = [PunchRecord("Alice", "u1", datetime.combine(DAY, value), "正常", 1) for value in (time(9, 30), time(18, 30))]
        result = OvertimeRule(config).calculate(day, punches)
        self.assertEqual(result.overtime_type, "休息日")
        self.assertEqual(result.overtime_minutes, 450)
        self.assertEqual(result.company_holiday_overtime_minutes, 150)

    def test_non_workday_switch_only_disables_company_part(self):
        workday = WorkdayConfig(count_non_workday_overtime=False)
        _, result = calculate("上午", "09:30", "22:00", workday=workday)
        self.assertEqual(result.overtime_minutes, 180)
        self.assertEqual(result.company_holiday_overtime_minutes, 0)
        self.assertEqual(result.approval_review_reason, "")
        _, result = calculate("全天", "09:30", "22:00", workday=workday)
        self.assertEqual(result.overtime_minutes, 0)
        self.assertEqual(result.approval_review_reason, "")

    def test_non_workday_dinner_setting_applies_to_afternoon_holiday(self):
        _, result = calculate("下午", "09:30", "22:00", rest_day=RestDayConfig(dinner_break_minutes=30))
        self.assertEqual(result.overtime_minutes, 480)
        self.assertEqual(result.dinner_break_minutes, 30)


class CompanyHolidayReportTest(unittest.TestCase):
    def test_full_holiday_additional_daily_cards_are_not_marked_exempt(self):
        config = calendar_config("全天")
        for punches, approval, expected_status in (
            ((("09:30", "正常"), ("18:30", "正常")), "加班10-19 09:30到10-19 18:30", "有效"),
            ((("09:30", "正常"),), "", "待补卡，不计入"),
            ((("09:30", "打卡无效"),), "", "待补卡，不计入"),
        ):
            with self.subTest(punches=punches), tempfile.TemporaryDirectory() as directory:
                day = DailyAttendance(
                    "Alice", "u1", DAY, "A 09:30-18:30",
                    additional_punches=punches, approval_text=approval,
                )
                result = OvertimeRule(config).calculate(day, [])
                path = Path(directory) / "report.xlsx"
                OvertimeReportWriter(config).write(path, [result])
                workbook = load_workbook(path, data_only=True)
                try:
                    detail = {cell.value: workbook["Alice"].cell(2, cell.column).value for cell in workbook["Alice"][1]}
                    self.assertEqual(detail["统计状态"], expected_status)
                    if result.punch_pair_complete:
                        self.assertEqual(result.overtime_minutes, 450)
                    else:
                        self.assertNotIn("无实际加班", detail["加班计算说明"])
                        self.assertEqual(workbook["待补卡通知"]["A2"].value, "Alice")
                finally:
                    workbook.close()

    def test_split_totals_exemptions_and_calendar_are_visible_and_consistent(self):
        config, result = calculate("上午", "09:30", "22:00", "加班10-19 09:30到10-19 12:00")
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "report.xlsx"
            OvertimeReportWriter(config).write(path, [result])
            workbook = load_workbook(path, data_only=True)
            try:
                overview = {cell.value: workbook["总览"].cell(2, cell.column).value for cell in workbook["总览"][1]}
                self.assertEqual(overview["工作日加班时间（分钟）"], 180)
                self.assertEqual(overview["休息日加班时间（分钟）"], 150)
                self.assertEqual(overview["加班时间（分钟）"], 330)
                self.assertEqual(overview["公司放假加班时间（分钟）"], 150)
                detail = {cell.value: workbook["Alice"].cell(2, cell.column).value for cell in workbook["Alice"][1]}
                self.assertEqual(detail["公司放假时段"], "上午")
                self.assertEqual(detail["公司放假说明"], "公司活动")
                self.assertEqual(detail["公司放假加班分钟"], 150)
                self.assertEqual(detail["当天应出勤分钟"], 300)
                daily = {cell.value: workbook["每日汇总"].cell(2, cell.column).value for cell in workbook["每日汇总"][1]}
                self.assertEqual(daily["工作日加班分钟"], 180)
                self.assertEqual(daily["休息日加班分钟"], 150)
                settings = dict(workbook["配置"].iter_rows(min_row=2, values_only=True))
                self.assertEqual(settings["公司放假 2026-10-19"], "上午；公司活动")
            finally:
                workbook.close()

    def test_full_day_without_punches_is_valid_in_report(self):
        config, result = calculate("全天")
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "report.xlsx"
            OvertimeReportWriter(config).write(path, [result])
            workbook = load_workbook(path, data_only=True)
            try:
                detail = {cell.value: workbook["Alice"].cell(2, cell.column).value for cell in workbook["Alice"][1]}
                self.assertEqual(detail["统计状态"], "公司放假，免考勤")
                self.assertIsNone(detail["漏卡"])
                self.assertEqual(workbook["待补卡通知"]["A2"].value, "未发现待补卡记录")
            finally:
                workbook.close()

    def test_missing_daily_row_on_company_full_day_does_not_claim_missing_punch(self):
        config = calendar_config("全天")
        results = [
            OvertimeRule(config).calculate(DailyAttendance("Alice", "u1", day, "休息"), [])
            for day in (date(2026, 10, 18), date(2026, 10, 20))
        ]
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "report.xlsx"
            OvertimeReportWriter(config).write(path, results)
            workbook = load_workbook(path, data_only=True)
            try:
                detail = {cell.value: workbook["Alice"].cell(3, cell.column).value for cell in workbook["Alice"][1]}
                self.assertEqual(detail["日类型"], "公司放假（全天）")
                self.assertIsNone(detail["漏卡"])
                self.assertEqual(detail["统计状态"], "公司放假，免考勤")
                self.assertEqual(detail["公司放假时段"], "全天")
            finally:
                workbook.close()
