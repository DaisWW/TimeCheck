from datetime import date
import unittest
from unittest.mock import patch

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from attendance_analyzer.config import AppConfig
from attendance_analyzer.excel_reader import DingTalkWorkbookReader


class DingTalkWorkbookReaderTest(unittest.TestCase):
    def test_dimension_scans_do_not_grow_with_attendance_rows(self):
        for daily in (True, False):
            scan_counts = []
            for row_count in (10, 100):
                workbook = Workbook()
                sheet = workbook.active
                sheet.append([])
                sheet.append([])
                if daily:
                    sheet.append([
                        "姓名", "UserId", "日期", "班次", "上班1打卡时间", "上班1打卡结果",
                        "下班1打卡时间", "下班1打卡结果", "关联的审批单",
                    ])
                    sheet.append([])
                    row = ["Alice", "u1", "2026-08-03", "A 09:30-18:30", "09:30", "正常", "20:00", "正常", ""]
                else:
                    sheet.append(["姓名", "UserId", "考勤日期", "打卡时间", "打卡结果"])
                    row = ["Alice", "u1", "2026-08-03", "2026-08-03 09:30", "正常"]
                for _ in range(row_count):
                    sheet.append(row)

                scans = []
                original = Worksheet.max_column.fget

                def max_column(worksheet):
                    scans.append(1)
                    return original(worksheet)

                reader = DingTalkWorkbookReader(AppConfig())
                with patch.object(Worksheet, "max_column", property(max_column)):
                    if daily:
                        self.assertEqual(len(reader._read_daily_records(sheet)), row_count)
                    else:
                        punches = reader._read_raw_punches(sheet)
                        self.assertEqual(len(punches[("u1", date(2026, 8, 3))]), row_count)
                scan_counts.append(len(scans))
                workbook.close()
            with self.subTest(daily=daily):
                self.assertEqual(scan_counts[0], scan_counts[1])

    def test_raw_punches_use_attendance_date_instead_of_clock_date(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([])
        sheet.append([])
        sheet.append(["姓名", "UserId", "考勤日期", "打卡时间", "打卡结果"])
        sheet.append(["Alice", "u1", "26-05-06 星期三", "2026-05-06 10:08", "正常"])
        sheet.append(["Alice", "u1", "26-05-06 星期三", "2026-05-07 01:24", "正常"])

        punches = DingTalkWorkbookReader(AppConfig())._read_raw_punches(sheet)

        self.assertEqual(
            [punch.occurred_at.isoformat(sep=" ") for punch in punches[("u1", date(2026, 5, 6))]],
            ["2026-05-06 10:08:00", "2026-05-07 01:24:00"],
        )

    def test_unparseable_raw_attendance_date_is_exposed_and_not_reassigned(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "原始记录"
        sheet.append([])
        sheet.append([])
        sheet.append(["姓名", "UserId", "考勤日期", "打卡时间", "打卡结果"])
        sheet.append(["Alice", "u1", "2026-02-30", "2026-03-01 10:08", "正常"])

        issues = []
        punches = DingTalkWorkbookReader(AppConfig())._read_raw_punches(sheet, issues)

        self.assertEqual(punches, {})
        self.assertEqual(len(issues), 1)
        self.assertEqual(issues[0].source_row, 4)
        self.assertEqual(issues[0].field_name, "考勤日期")
        self.assertEqual(issues[0].raw_value, "2026-02-30")

    def test_unparseable_daily_punch_time_is_recorded_for_manual_review(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "每日统计"
        sheet.append([])
        sheet.append([])
        sheet.append(
            [
                "姓名",
                "UserId",
                "日期",
                "班次",
                "上班1打卡时间",
                "上班1打卡结果",
                "下班1打卡时间",
                "下班1打卡结果",
                "关联的审批单",
            ]
        )
        sheet.append([])
        sheet.append(["Alice", "u1", "2026-07-01", "A 09:30-18:30", "25:00", "", "19:00", "", ""])

        issues = []
        records = DingTalkWorkbookReader(AppConfig())._read_daily_records(sheet, issues)

        self.assertEqual(len(records), 1)
        self.assertEqual(len(issues), 1)
        self.assertEqual(issues[0].source_row, 5)
        self.assertEqual(issues[0].field_name, "上班1打卡时间")


if __name__ == "__main__":
    unittest.main()
