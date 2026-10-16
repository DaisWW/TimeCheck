from datetime import date, datetime
from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook, load_workbook

from attendance_analyzer.analyzer import AttendanceAnalyzer
from attendance_analyzer.config import AppConfig
from attendance_analyzer.excel_reader import DingTalkWorkbookReader, WorkbookInputError
from attendance_analyzer.report_writer import OvertimeReportWriter
from attendance_analyzer.time_utils import combine_date_time, parse_date_text, parse_datetime, parse_time


def valid_workbook():
    workbook = Workbook()
    daily = workbook.active
    daily.title = "每日统计"
    daily.append([])
    daily.append([])
    daily.append([
        "姓名", "UserId", "日期", "班次", "上班1打卡时间", "上班1打卡结果",
        "下班1打卡时间", "下班1打卡结果", "关联的审批单",
    ])
    daily.append([])
    daily.append(["Alice", "u1", "26-08-03 星期一", "A 09:30-18:30", "09:30", "正常", "20:00", "正常", ""])
    raw = workbook.create_sheet("原始记录")
    raw.append([])
    raw.append([])
    raw.append(["姓名", "UserId", "考勤日期", "打卡时间", "打卡结果"])
    raw.append(["Alice", "u1", "26-08-03 星期一", "2026-08-03 09:30", "正常"])
    raw.append(["Alice", "u1", "26-08-03 星期一", "2026-08-03 20:00", "正常"])
    return workbook


class InputValidationTest(unittest.TestCase):
    def read(self, workbook):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "attendance.xlsx"
            workbook.save(path)
            try:
                return DingTalkWorkbookReader(AppConfig()).read(path)
            finally:
                workbook.close()

    def test_valid_export_and_trailing_blank_rows_are_accepted(self):
        workbook = valid_workbook()
        workbook["每日统计"].cell(7, 1).number_format = "@"
        workbook["原始记录"].cell(8, 1).number_format = "@"
        data = self.read(workbook)
        self.assertEqual(len(data.daily_records), 1)
        self.assertEqual(len(data.raw_punches[("u1", date(2026, 8, 3))]), 2)
        self.assertEqual(data.input_issues, [])

    def test_all_bad_cells_are_reported_with_source_and_original_value(self):
        workbook = valid_workbook()
        workbook["每日统计"]["E5"] = "10:30 trailing"
        workbook["每日统计"]["I5"] = "产假08-03 09:00到08-03 18:00"
        workbook["原始记录"]["D4"] = "2026-08-03 25:00"
        with self.assertRaises(WorkbookInputError) as raised:
            self.read(workbook)
        error = str(raised.exception)
        for fragment in ("每日统计 第 5 行", "原始记录 第 4 行", "10:30 trailing", "产假", "25:00"):
            self.assertIn(fragment, error)
        self.assertGreaterEqual(len(raised.exception.issues), 3)

    def test_missing_or_duplicate_required_header_is_error(self):
        for value in (None, "姓名"):
            with self.subTest(value=value):
                workbook = valid_workbook()
                workbook["每日统计"]["I3"] = value
                with self.assertRaisesRegex(WorkbookInputError, "表头"):
                    self.read(workbook)

    def test_missing_names_are_reported_in_both_sheets(self):
        workbook = valid_workbook()
        workbook["每日统计"]["A5"] = None
        workbook["原始记录"]["A4"] = None
        with self.assertRaises(WorkbookInputError) as raised:
            self.read(workbook)
        for fragment in ("每日统计 第 5 行 [姓名]", "原始记录 第 4 行 [姓名]", "缺少姓名"):
            self.assertIn(fragment, str(raised.exception))

    def test_duplicate_day_and_orphan_raw_record_are_errors(self):
        workbook = valid_workbook()
        daily = workbook["每日统计"]
        daily.append([cell.value for cell in daily[5]])
        workbook["原始记录"].append(["Bob", "u2", "26-08-03 星期一", "2026-08-03 12:00", "正常"])
        with self.assertRaises(WorkbookInputError) as raised:
            self.read(workbook)
        self.assertIn("重复", str(raised.exception))
        self.assertIn("找不到对应", str(raised.exception))

    def test_numeric_identifier_unknown_status_and_formula_are_errors(self):
        workbook = valid_workbook()
        workbook["每日统计"]["B5"] = 12345678901234567890
        workbook["每日统计"]["E5"] = '=TEXT(TIME(9,30,0),"hh:mm")'
        workbook["原始记录"]["E4"] = "位置异常但不知道是否有效"
        with self.assertRaises(WorkbookInputError) as raised:
            self.read(workbook)
        for fragment in ("文本保存", "公式", "未知或空的打卡结果"):
            self.assertIn(fragment, str(raised.exception))

    def test_reversed_fallback_and_unrelated_raw_date_are_errors(self):
        workbook = valid_workbook()
        workbook["每日统计"]["E5"] = "18:00"
        workbook["每日统计"]["G5"] = "09:00"
        workbook["原始记录"]["D5"] = "2026-08-05 20:00"
        with self.assertRaises(WorkbookInputError) as raised:
            self.read(workbook)
        self.assertIn("先后顺序", str(raised.exception))
        self.assertIn("考勤当天或次日", str(raised.exception))

    def test_date_and_time_must_match_the_entire_value(self):
        for value in ("garbage 2026-08-03", "2026-08-03 or 2026-08-04", "2026-08-03 星期五", "2026-02-30"):
            with self.subTest(value=value):
                self.assertIsNone(parse_date_text(value))
        for value in ("garbage 10:30", "2026-08-04 10:30", "10:30:99", "10:30 trailing"):
            with self.subTest(value=value):
                self.assertIsNone(combine_date_time(date(2026, 8, 3), value))
        self.assertIsNone(parse_datetime("2026-08-03 10:3"))
        self.assertEqual(combine_date_time(date(2026, 8, 3), "次日 01:24:59"), datetime(2026, 8, 4, 1, 24, 59))
        with self.assertRaises(ValueError):
            parse_time("09:30 trailing")

    def test_config_rejects_coercion_and_unknown_keys(self):
        for section, key, value in (("workday", "count_non_workday_overtime", "随便"), ("overtime", "dinner_break_minutes", 2.5), ("workday", "work_minute", 450)):
            with self.subTest(key=key):
                data = AppConfig().to_dict()
                data[section][key] = value
                with self.assertRaises(ValueError):
                    AppConfig.from_dict(data)


class InputIssueResolutionTest(unittest.TestCase):
    def analyze(self, workbook):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "attendance.xlsx"
            workbook.save(path)
            try:
                analyzer = AttendanceAnalyzer(AppConfig())
                results = analyzer.analyze_overtime(path)
                return results, analyzer.last_input_issues
            finally:
                workbook.close()

    def test_updated_old_punches_are_audited_without_pending_status(self):
        workbook = valid_workbook()
        raw = workbook["原始记录"]
        for status in ("打卡无效:此记录已被更新", "打卡无效：此记录已被更新"):
            raw.append(["Alice", "u1", "26-08-03 星期一", "2026-08-03 19:45", status])
        results, issues = self.analyze(workbook)
        result = results[0]
        self.assertEqual(issues, [])
        self.assertEqual(result.overtime_minutes, 60)
        self.assertEqual(result.last_clock_out, datetime(2026, 8, 3, 20))
        self.assertEqual(result.valid_raw_punch_count, 2)
        self.assertIn("原始记录第 6、7 行", "；".join(result.notes))
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "report.xlsx"
            OvertimeReportWriter(AppConfig()).write(path, results, issues)
            book = load_workbook(path, data_only=True)
            try:
                row = {cell.value: book["Alice"].cell(2, cell.column).value for cell in book["Alice"][1]}
                self.assertEqual(row["人工确认条数"], 0)
                self.assertEqual(row["人工确认状态"], "无")
                self.assertEqual(row["异常状态"], "正常")
                self.assertEqual(row["统计状态"], "有效")
                self.assertIn("已自动排除", row["备注"])
            finally:
                book.close()

    def test_incomplete_punch_pair_keeps_old_punch_warning(self):
        workbook = valid_workbook()
        workbook["每日统计"]["G5"] = None
        workbook["原始记录"]["E5"] = "打卡无效:此记录已被更新"
        results, issues = self.analyze(workbook)
        self.assertFalse(results[0].punch_pair_complete)
        self.assertEqual(results[0].overtime_minutes, 0)
        self.assertEqual(len(issues), 1)
        self.assertEqual(issues[0].issue_type, "原始数据被规则排除")

    def test_other_invalid_punch_reasons_still_require_confirmation(self):
        workbook = valid_workbook()
        workbook["原始记录"].append([
            "Alice", "u1", "26-08-03 星期一", "2026-08-03 19:45", "打卡无效:位置异常",
        ])
        results, issues = self.analyze(workbook)
        self.assertTrue(results[0].punch_pair_complete)
        self.assertEqual(len(issues), 1)
        self.assertEqual(issues[0].raw_value, "打卡无效:位置异常")

    def test_approved_supplement_matching_the_selected_card_is_resolved(self):
        for raw_time in ("2026-08-03 20:00", "2026-08-03 20:00:59"):
            with self.subTest(raw_time=raw_time):
                workbook = valid_workbook()
                workbook["每日统计"]["I5"] = "补卡申请08-03 20:00到08-03 20:00"
                workbook["原始记录"]["D5"] = raw_time
                workbook["原始记录"]["E5"] = "补卡审批通过"
                results, issues = self.analyze(workbook)
                self.assertEqual(issues, [])
                self.assertEqual(results[0].overtime_minutes, 60)
                self.assertIn("2026-08-03 20:00", "；".join(results[0].notes))
                self.assertIn("补卡申请已自动核对", "；".join(results[0].notes))

    def test_unmatched_supplement_time_or_status_remains_pending(self):
        for time_text, status in (("2026-08-03 20:00", "正常"), ("2026-08-03 19:59", "补卡审批通过")):
            with self.subTest(time_text=time_text, status=status):
                workbook = valid_workbook()
                workbook["每日统计"]["I5"] = "补卡申请08-03 20:00到08-03 20:00"
                workbook["原始记录"]["D5"] = time_text
                workbook["原始记录"]["E5"] = status
                results, issues = self.analyze(workbook)
                self.assertTrue(results[0].punch_pair_complete)
                self.assertEqual(len(issues), 1)
                self.assertIn("未匹配", issues[0].message)
                self.assertFalse(issues[0].blocks_summary)

    def test_approved_single_punch_does_not_replace_the_missing_pair(self):
        workbook = valid_workbook()
        workbook["每日统计"]["E5"] = None
        workbook["每日统计"]["I5"] = "补卡申请08-03 20:00到08-03 20:00"
        workbook["原始记录"].delete_rows(4)
        workbook["原始记录"]["E4"] = "补卡审批通过"
        results, issues = self.analyze(workbook)
        self.assertFalse(results[0].punch_pair_complete)
        self.assertEqual(results[0].overtime_minutes, 0)
        self.assertEqual(len(issues), 1)

    def test_only_the_selected_fallback_can_resolve_a_supplement(self):
        for use_fallback in (True, False):
            with self.subTest(use_fallback=use_fallback):
                workbook = valid_workbook()
                workbook["每日统计"]["H5"] = "补卡审批通过"
                workbook["每日统计"]["I5"] = "补卡申请08-03 20:00到08-03 20:00"
                if use_fallback:
                    workbook["原始记录"].delete_rows(4, 2)
                results, issues = self.analyze(workbook)
                self.assertEqual(results[0].used_daily_fallback, use_fallback)
                self.assertEqual(len(issues), 0 if use_fallback else 1)

    def test_partial_supplement_matches_keep_the_unmatched_request(self):
        workbook = valid_workbook()
        workbook["每日统计"]["I5"] = (
            "补卡申请08-03 09:30到08-03 09:30，补卡申请08-03 20:00到08-03 20:00"
        )
        workbook["原始记录"]["E4"] = "补卡审批通过"
        results, issues = self.analyze(workbook)
        self.assertEqual(len(issues), 1)
        self.assertIn("20:00", issues[0].message)
        self.assertNotIn("09:30", issues[0].message)
        self.assertIn("09:30", "；".join(results[0].notes))

    def test_another_employee_or_attendance_day_cannot_resolve_the_request(self):
        for other_employee in (True, False):
            with self.subTest(other_employee=other_employee):
                workbook = valid_workbook()
                workbook["每日统计"]["I5"] = "补卡申请08-03 20:00到08-03 20:00"
                other = [cell.value for cell in workbook["每日统计"][5]]
                other[8] = ""
                if other_employee:
                    other[0], other[1] = "Bob", "u2"
                    attendance_date, clock_date = "26-08-03 星期一", "2026-08-03"
                else:
                    other[2] = "26-08-04 星期二"
                    attendance_date, clock_date = "26-08-04 星期二", "2026-08-04"
                workbook["每日统计"].append(other)
                workbook["原始记录"].append([other[0], other[1], attendance_date, f"{clock_date} 09:30", "正常"])
                workbook["原始记录"].append([other[0], other[1], attendance_date, f"{clock_date} 20:00", "补卡审批通过"])
                _, issues = self.analyze(workbook)
                self.assertEqual(len(issues), 1)
                self.assertEqual(issues[0].employee_name, "Alice")
                self.assertEqual(issues[0].work_date, date(2026, 8, 3))

    def test_next_day_approved_card_matches_its_original_attendance_day(self):
        workbook = valid_workbook()
        workbook["每日统计"]["G5"] = "次日 01:00"
        workbook["每日统计"]["I5"] = "补卡申请08-04 01:00到08-04 01:00"
        workbook["原始记录"]["D5"] = "2026-08-04 01:00"
        workbook["原始记录"]["E5"] = "补卡审批通过"
        results, issues = self.analyze(workbook)
        self.assertEqual(issues, [])
        self.assertEqual(results[0].work_date, date(2026, 8, 3))
        self.assertEqual(results[0].last_clock_out, datetime(2026, 8, 4, 1))

    def test_non_point_supplement_interval_is_not_automatically_resolved(self):
        workbook = valid_workbook()
        workbook["每日统计"]["I5"] = "补卡申请08-03 09:30到08-03 20:00"
        workbook["原始记录"]["E4"] = "补卡审批通过"
        workbook["原始记录"]["E5"] = "补卡审批通过"
        _, issues = self.analyze(workbook)
        self.assertEqual(len(issues), 1)

    def test_bad_approval_or_unknown_punch_is_still_blocking(self):
        for bad_approval, bad_punch in ((True, False), (False, True)):
            with self.subTest(bad_approval=bad_approval):
                workbook = valid_workbook()
                workbook["每日统计"]["I5"] = (
                    "补卡申请08-03 20:00到08-03 20:00，产假格式未知"
                    if bad_approval else "补卡申请08-03 20:00到08-03 20:00"
                )
                workbook["原始记录"]["E5"] = "补卡审批通过"
                if bad_punch:
                    workbook["原始记录"].append([
                        "Alice", "u1", "26-08-03 星期一", "2026-08-03 19:45", "未知状态",
                    ])
                with self.assertRaises(WorkbookInputError):
                    self.analyze(workbook)
